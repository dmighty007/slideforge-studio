import io
import math
import logging
import base64
import re
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

from slideforge.bridge.pptx_animations import apply_slide_animations

logger = logging.getLogger(__name__)

class DesignSystem:
    # Default design constants inspired by the reference script
    DEFAULT_FONT = "Arial"
    ACCENT_COLOR = RGBColor(0, 51, 102)  # Dark Blue
    TEXT_COLOR = RGBColor(44, 62, 80)
    BG_COLOR = RGBColor(255, 255, 255)


class HtmlTextRunParser(HTMLParser):
    BLOCK_TAGS = {"div", "p", "li", "section", "article", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.runs: list[dict[str, Any]] = []
        self._style_stack: list[dict[str, Any]] = [{"bold": False, "italic": False, "underline": False}]
        self._tag_stack: list[tuple[str, bool]] = []

    @property
    def current_style(self) -> dict[str, Any]:
        return self._style_stack[-1]

    def _push(self, **updates):
        self._style_stack.append({**self.current_style, **updates})

    def _mark_tag(self, tag: str, pushed: bool):
        self._tag_stack.append((tag, pushed))

    def _pop_tag(self, tag: str) -> bool:
        if not self._tag_stack:
            return False
        if self._tag_stack[-1][0] == tag:
            _, pushed = self._tag_stack.pop()
        else:
            pushed = False
            for idx in range(len(self._tag_stack) - 1, -1, -1):
                if self._tag_stack[idx][0] == tag:
                    _, pushed = self._tag_stack.pop(idx)
                    break
        if pushed and len(self._style_stack) > 1:
            self._style_stack.pop()
        return pushed

    def _newline(self):
        if self.runs and self.runs[-1].get("text") == "\n":
            return
        self.runs.append({"text": "\n", **self.current_style})

    def _parse_inline_style(self, style_text: str) -> dict[str, Any]:
        updates: dict[str, Any] = {}
        for declaration in str(style_text or "").split(";"):
            if ":" not in declaration:
                continue
            prop, value = declaration.split(":", 1)
            prop = prop.strip().lower()
            value = value.strip()
            if not value:
                continue
            if prop == "font-size":
                updates["fontSize"] = value
            elif prop == "font-family":
                updates["fontFamily"] = value
            elif prop == "color":
                updates["color"] = value
            elif prop == "font-weight":
                updates["fontWeight"] = value
                if value.lower() == "bold" or re.search(r"\d+", value) and int(re.search(r"\d+", value).group(0)) >= 600:
                    updates["bold"] = True
            elif prop == "font-style":
                updates["fontStyle"] = value
                if value.lower() == "italic":
                    updates["italic"] = True
            elif prop == "text-decoration" and "underline" in value.lower():
                updates["underline"] = True
            elif prop == "vertical-align":
                if value.lower() in {"super", "sup", "sub"}:
                    updates["verticalAlign"] = value.lower()
        return updates

    def _attrs_to_updates(self, attrs) -> dict[str, Any]:
        attrs_map = {str(key).lower(): value for key, value in attrs}
        updates = self._parse_inline_style(attrs_map.get("style", ""))
        if attrs_map.get("face"):
            updates["fontFamily"] = attrs_map["face"]
        if attrs_map.get("color"):
            updates["color"] = attrs_map["color"]
        if attrs_map.get("size"):
            # Legacy <font size="1..7"> values are approximate browser defaults.
            size_map = {"1": "10px", "2": "13px", "3": "16px", "4": "18px", "5": "24px", "6": "32px", "7": "48px"}
            updates["fontSize"] = size_map.get(str(attrs_map["size"]).strip(), attrs_map["size"])
        return updates

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        attr_updates = self._attrs_to_updates(attrs)
        if tag in {"strong", "b"}:
            self._push(**attr_updates, bold=True)
            self._mark_tag(tag, True)
        elif tag in {"em", "i"}:
            self._push(**attr_updates, italic=True)
            self._mark_tag(tag, True)
        elif tag == "u":
            self._push(**attr_updates, underline=True)
            self._mark_tag(tag, True)
        elif tag == "sup":
            self._push(**attr_updates, verticalAlign="super")
            self._mark_tag(tag, True)
        elif tag == "sub":
            self._push(**attr_updates, verticalAlign="sub")
            self._mark_tag(tag, True)
        elif tag == "a":
            href = str(dict(attrs).get("href") or "").strip()
            link = {"link": href} if re.match(r"(https?:|mailto:)", href, re.I) else {}
            self._push(**attr_updates, **link)
            self._mark_tag(tag, True)
        elif tag == "br":
            self._newline()
        elif tag in {"span", "font"}:
            self._push(**attr_updates)
            self._mark_tag(tag, True)
        elif tag in self.BLOCK_TAGS:
            if self.runs:
                self._newline()
            if attr_updates:
                self._push(**attr_updates)
                self._mark_tag(tag, True)
            else:
                self._mark_tag(tag, False)
        elif attr_updates:
            self._push(**attr_updates)
            self._mark_tag(tag, True)
        else:
            self._mark_tag(tag, False)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in {"strong", "b", "em", "i", "u", "span", "font", "sup", "sub", "a"}:
            self._pop_tag(tag)
        elif tag in self.BLOCK_TAGS:
            self._pop_tag(tag)
            self._newline()
        else:
            self._pop_tag(tag)

    def handle_data(self, data):
        if data:
            self.runs.append({"text": unescape(data), **self.current_style})


class HtmlListParser(HTMLParser):
    """Splits content that is one HTML list (<ol>/<ul>, possibly nested) into items: html, level, kind, style."""

    def __init__(self):
        super().__init__(convert_charrefs=False)
        self.items: list[dict[str, Any]] = []
        self.lists: list[tuple[str, str]] = []  # (tag, list-style-type)
        self.outside_text = ""
        self._current: Optional[dict[str, Any]] = None

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in {"ol", "ul"}:
            style = dict(attrs).get("style") or ""
            match = re.search(r"list-style-type\s*:\s*([a-z-]+)", style, re.I)
            inherited = self.lists[-1][1] if self.lists and self.lists[-1][0] == tag else ""
            self.lists.append((tag, match.group(1).lower() if match else inherited))
            self._current = None
        elif tag == "li" and self.lists:
            list_tag, list_style = self.lists[-1]
            self._current = {
                "html": "",
                "level": len(self.lists) - 1,
                "kind": "numbered" if list_tag == "ol" else "bullet",
                "numberStyle": list_style or "decimal",
            }
            self.items.append(self._current)
        elif self._current is not None:
            self._current["html"] += self.get_starttag_text() or ""

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in {"ol", "ul"}:
            if self.lists:
                self.lists.pop()
            self._current = None
        elif tag == "li":
            self._current = None
        elif self._current is not None:
            self._current["html"] += f"</{tag}>"

    def _text(self, text: str):
        if self._current is not None:
            self._current["html"] += text
        elif not self.lists:
            self.outside_text += text

    def handle_data(self, data):
        self._text(data)

    def handle_entityref(self, name):
        self._text(f"&{name};")

    def handle_charref(self, name):
        self._text(f"&#{name};")


def html_list_items(content: Any) -> Optional[list[dict[str, Any]]]:
    """The items of content that is a single HTML list, or None for any other content."""
    text = str(content or "")
    if not re.search(r"<(ol|ul)[\s>]", text, re.I):
        return None
    parser = HtmlListParser()
    parser.feed(text)
    if parser.outside_text.strip() or not parser.items:
        return None
    return [item for item in parser.items if unescape(re.sub(r"<[^>]+>", "", item["html"])).strip()]


def html_to_text_runs(content: Any) -> list[dict[str, Any]]:
    parser = HtmlTextRunParser()
    parser.feed(str(content or ""))
    runs = []
    for run in parser.runs:
        text = run.get("text", "")
        if text == "\n":
            if runs and runs[-1].get("text") == "\n":
                continue
            runs.append(run)
        elif text:
            runs.append(run)
    while runs and runs[-1].get("text") == "\n":
        runs.pop()
    return runs or [{"text": ""}]

class PPTXExporter:
    def __init__(self, state: Dict[str, Any], project_root: Optional[Path | str] = None):
        self.state = state
        self.prs = Presentation()
        self.project_root = self._resolve_project_root(project_root)

        # Mapping from SlideForge page setup to PPTX dimensions
        # The editor's page sizes (frontend/js/core/pageSetup.js), in canvas pixels at 96 per inch. Only the two
        # legacy widescreen names were known here, so "Talk 16:9" and "Lecture 16:10" decks came out as 4:3.
        setup = str(state.get("pageSetup") or "standard-4-3")
        self.base_w, self.base_h = self.PAGE_SETUPS.get(setup, self.PAGE_SETUPS["standard-4-3"])
        self.prs.slide_width = Inches(self.base_w / 96)
        self.prs.slide_height = Inches(self.base_h / 96)
        self.theme_background_css = str(state.get("themeBackgroundCss") or "")
        self.theme = self._theme_defaults(str(state.get("presentationTheme") or "editorial"))
        self.slide_bg_rgb = self.theme["background"]

    PAGE_SETUPS = {
        "standard-4-3": (1024, 768),
        "talk-16-9": (1280, 720),
        "widescreen-16-9": (1280, 720),
        "lecture-16-10": (1280, 800),
        "widescreen-16-10": (1280, 800),
        "paper-letter": (816, 1056),
        "poster-portrait": (1440, 1920),
    }

    def _theme_gradient(self) -> Optional[tuple[RGBColor, RGBColor, float]]:
        """The first and last colour of the theme's base slide gradient and its PowerPoint angle, if it has one.

        Any angle: Graphite, Circuit, Afterglow and Midnight Garden use diagonal (155-160deg) gradients, which were
        skipped, so their slides came out flat. CSS angles turn clockwise from "towards the top"; PowerPoint's turn
        counter-clockwise from "left to right", so PowerPoint = 90 - CSS (180deg, top to bottom, is 270).
        """
        pattern = r"linear-gradient\(\s*(-?[\d.]+)deg\s*,([^()]*(?:\([^()]*\)[^()]*)*)\)"
        for angle, gradient in reversed(re.findall(pattern, self.theme_background_css)):
            colours = re.findall(r"#[0-9a-fA-F]{6}\b", gradient)
            if len(colours) >= 2 and colours[0].lower() != colours[-1].lower():
                return (
                    self._parse_color(colours[0], self.slide_bg_rgb),
                    self._parse_color(colours[-1], self.slide_bg_rgb),
                    (90 - float(angle)) % 360,
                )
        return None

    def _theme_defaults(self, theme_id: str) -> dict[str, RGBColor]:
        themes = {
            "editorial": ("#EEF2F7", "#2E2E2E", "#6B7280", "#3B82F6", "#A7C7E7"),
            "blueprint": ("#F1F7FD", "#10233B", "#5F7287", "#2563EB", "#0F766E"),
            "fieldnotes": ("#EFE4D0", "#2F261D", "#6F6559", "#7A8F47", "#B48A4A"),
            "monograph": ("#ECEFF3", "#171717", "#6B7280", "#1F2937", "#9CA3AF"),
            "graphite": ("#111827", "#F5F7FB", "#94A3B8", "#22D3EE", "#38BDF8"),
            "horizon": ("#0B2442", "#EEF4FF", "#AFBDD6", "#7DD3FC", "#4F7CFF"),
            "chalkboard": ("#10261F", "#F8F3E7", "#D4CBB7", "#F5D76E", "#8ED1C7"),
            "circuit": ("#0A2320", "#ECF7F5", "#9AB8B3", "#63E6D8", "#1FB6A6"),
            "afterglow": ("#1A2034", "#F5F7FF", "#BDC4DE", "#F3B76A", "#6E86FF"),
            "sage": ("#EDF4E8", "#25342D", "#68786E", "#4F7D5A", "#B7A66C"),
            "porcelain": ("#EEF7F9", "#233047", "#6C7890", "#477CA4", "#C08393"),
            "rosewater": ("#F7E7E4", "#3A2B32", "#7F6B72", "#B76E79", "#AA8650"),
            "buttercup": ("#F5EBC5", "#332B15", "#756C4C", "#D59D20", "#5F9F76"),
            "tidepool": ("#E4F4F3", "#17323A", "#5F7880", "#168B8B", "#4B8FC5"),
            "lavender": ("#EFEAF8", "#302C45", "#756F8D", "#7D6BB3", "#CF8A72"),
            "midnightGarden": ("#10231A", "#EFF7ED", "#B7C9B7", "#9AE6B4", "#5FAF79"),
            "retroPop": ("#FFE7DF", "#28212A", "#745F76", "#EF476F", "#06D6A0"),
        }
        bg, text, muted, accent, accent2 = themes.get(theme_id, themes["editorial"])
        return {
            "background": self._parse_color(bg, DesignSystem.BG_COLOR),
            "text": self._parse_color(text, DesignSystem.TEXT_COLOR),
            "muted": self._parse_color(muted, RGBColor(100, 116, 139)),
            "accent": self._parse_color(accent, DesignSystem.ACCENT_COLOR),
            "accent2": self._parse_color(accent2, DesignSystem.ACCENT_COLOR),
        }

    def _resolve_project_root(self, project_root: Optional[Path | str] = None) -> Path:
        if project_root:
            return Path(project_root).expanduser().resolve()
        try:
            from django.conf import settings

            base_dir = getattr(settings, "BASE_DIR", None)
            if base_dir:
                return Path(base_dir).expanduser().resolve()
        except Exception:
            pass
        return Path(__file__).resolve().parent.parent

    def _safe_asset_roots(self) -> list[Path]:
        roots = [
            self.project_root / "media",
            self.project_root / "static",
            self.project_root / "staticfiles",
            self.project_root / "assets",
        ]
        try:
            from django.conf import settings

            for setting_name in ("MEDIA_ROOT", "STATIC_ROOT"):
                configured = getattr(settings, setting_name, None)
                if configured:
                    roots.append(Path(configured))
            for static_dir in getattr(settings, "STATICFILES_DIRS", []) or []:
                roots.append(Path(static_dir))
        except Exception:
            pass

        resolved = []
        seen = set()
        for root in roots:
            try:
                path = Path(root).expanduser().resolve()
            except Exception:
                continue
            if path in seen:
                continue
            seen.add(path)
            resolved.append(path)
        return resolved

    def _media_candidates(self, url_path: str) -> list[Path]:
        """Files behind a /media/... URL: the configured MEDIA_ROOT first, then <project_root>/media."""
        relative = url_path.removeprefix("/media/")
        candidates = []
        try:
            from django.conf import settings

            media_root = getattr(settings, "MEDIA_ROOT", None)
            if media_root:
                candidates.append(Path(media_root) / relative)
        except Exception:
            pass
        candidates.append(self.project_root / "media" / relative)
        return candidates

    def _resolve_safe_local_image_path(self, content: str) -> Optional[Path]:
        normalized = str(content or "").strip()
        if not normalized or "://" in normalized:
            return None

        candidates: list[Path] = []
        if normalized.startswith("/media/"):
            candidates.extend(self._media_candidates(normalized))
        elif normalized.startswith("/static/"):
            candidates.append(self.project_root / normalized.lstrip("/"))
            candidates.append(self.project_root / "staticfiles" / normalized.removeprefix("/static/"))
        elif normalized.startswith("/assets/"):
            candidates.append(self.project_root / normalized.lstrip("/"))
        elif normalized.startswith("/"):
            candidates.append(Path(normalized))
        else:
            candidates.append(self.project_root / normalized)

        safe_roots = self._safe_asset_roots()
        for candidate in candidates:
            try:
                path = candidate.expanduser().resolve()
            except Exception:
                continue
            if not path.exists() or not path.is_file():
                continue
            if any(path == root or path.is_relative_to(root) for root in safe_roots):
                return path
        return None

    def _parse_font_weight(self, value: Any) -> int:
        if isinstance(value, (int, float)):
            return int(value)
        normalized = str(value or "400").strip().lower()
        if normalized == "normal":
            return 400
        if normalized == "bold":
            return 700
        match = re.search(r"\d+", normalized)
        if match:
            return int(match.group(0))
        return 400

    def _px_to_inches_w(self, px: float) -> Inches:
        # Scale pixels relative to the logical canvas width
        return Inches((px / self.base_w) * (self.prs.slide_width / Inches(1)))

    def _px_to_pt(self, px: float) -> Pt:
        """A font size in canvas pixels as points on the PowerPoint slide. The canvas is laid out at 96 px per inch
        (1024 px across a 10.67 in slide), so 24 px is 18 pt; written as 24 pt, text came out a third too large."""
        inches_per_px = (self.prs.slide_width / Inches(1)) / self.base_w
        return Pt(max(1.0, round(px * inches_per_px * 72 * 2) / 2))

    def _px_to_inches_h(self, px: float) -> Inches:
        # Scale pixels relative to the logical canvas height
        return Inches((px / self.base_h) * (self.prs.slide_height / Inches(1)))

    def _parse_color_components(self, color_str: Optional[str]) -> Optional[tuple[int, int, int, float]]:
        if not color_str:
            return None
        normalized = str(color_str).strip()
        if not normalized or normalized.lower() in {"transparent", "none", "inherit", "currentcolor"}:
            return None
        try:
            if normalized.startswith("#"):
                h = normalized.lstrip("#")
                if len(h) == 3:
                    h = "".join([c*2 for c in h])
                if len(h) >= 6:
                    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16), 1.0
            elif normalized.startswith("rgb"):
                matches = re.findall(r"[\d.]+", normalized)
                if len(matches) >= 3:
                    alpha = float(matches[3]) if len(matches) >= 4 else 1.0
                    return (
                        max(0, min(255, int(float(matches[0])))),
                        max(0, min(255, int(float(matches[1])))),
                        max(0, min(255, int(float(matches[2])))),
                        max(0.0, min(1.0, alpha)),
                    )
        except Exception:
            return None
        return None

    def _composite_color(self, color: tuple[int, int, int, float], bg: Optional[RGBColor] = None) -> RGBColor:
        r, g, b, alpha = color
        if alpha >= 0.995:
            return RGBColor(r, g, b)
        bg = bg or self.slide_bg_rgb
        return RGBColor(
            round(r * alpha + bg[0] * (1 - alpha)),
            round(g * alpha + bg[1] * (1 - alpha)),
            round(b * alpha + bg[2] * (1 - alpha)),
        )

    def _parse_color(self, color_str: Optional[str], default: RGBColor = DesignSystem.TEXT_COLOR) -> RGBColor:
        parsed = self._parse_color_components(color_str)
        if parsed:
            return self._composite_color(parsed)
        return default

    def _is_transparent(self, color_str: Optional[str]) -> bool:
        normalized = str(color_str or "").strip().lower()
        if normalized in {"", "transparent", "none"}:
            return True
        parsed = self._parse_color_components(normalized)
        return bool(parsed and parsed[3] <= 0.01)

    def _parse_px(self, value: Any, fallback: float = 100) -> float:
        match = re.search(r"-?\d+(?:\.\d+)?", str(value or ""))
        if not match:
            return fallback
        try:
            return float(match.group(0))
        except ValueError:
            return fallback

    def _set_text_common(self, run: Any, styles: dict[str, Any], *, extra: Optional[dict[str, bool]] = None):
        extra = extra or {}
        font_size = int(self._parse_px(extra.get("fontSize") or styles.get("fontSize", "18"), 18))
        font_color = self._parse_color(extra.get("color") or styles.get("color"), self.theme["text"])
        font_name = str(extra.get("fontFamily") or styles.get("fontFamily", DesignSystem.DEFAULT_FONT)).split(",")[0].replace('"', "").strip()
        run.font.size = self._px_to_pt(font_size)
        run.font.color.rgb = font_color
        run.font.name = font_name
        run.font.bold = self._parse_font_weight(extra.get("fontWeight") or styles.get("fontWeight")) >= 600 or extra.get("bold", False)
        run.font.italic = (extra.get("fontStyle") or styles.get("fontStyle")) == "italic" or extra.get("italic", False)
        run.font.underline = bool(extra.get("underline", False))
        # "All caps" is a display setting in PowerPoint too: the text itself keeps its case.
        if str(styles.get("textTransform", "")).lower() == "uppercase":
            run.font._element.set("cap", "all")
        if extra.get("verticalAlign") == "super" and hasattr(run.font, "superscript"):
            run.font.superscript = True
        if extra.get("verticalAlign") == "sub" and hasattr(run.font, "subscript"):
            run.font.subscript = True
        if extra.get("link"):
            run.hyperlink.address = extra["link"]

    def _paragraph_alignment(self, styles: dict[str, Any]):
        align = str(styles.get("textAlign", "left")).lower()
        if align == "center":
            return PP_ALIGN.CENTER
        if align == "right":
            return PP_ALIGN.RIGHT
        if align == "justify":
            return PP_ALIGN.JUSTIFY
        return PP_ALIGN.LEFT

    def _get_image_stream(self, content: str) -> Optional[io.BytesIO]:
        if not content:
            return None

        if content.startswith("data:"):
            try:
                base64_data = content.split(",")[1]
                return io.BytesIO(base64.b64decode(base64_data))
            except Exception:
                return None

        try:
            path = self._resolve_safe_local_image_path(content)
            if path:
                with open(path, "rb") as f:
                    return io.BytesIO(f.read())
        except Exception as e:
            logger.warning("Could not resolve image path %s: %s", content, e)

        return None

    def _apply_fill(self, shape: Any, color: Optional[str], default: Optional[RGBColor] = None):
        if self._is_transparent(color):
            shape.fill.background()
            return
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._parse_color(color, default or self.theme["accent"])

    def _corner_radius_px(self, value: Any) -> float:
        """The first corner radius of a CSS border-radius ("10px", "10px 0 0 10px", "999px"), in pixels."""
        match = re.search(r"-?\d+(?:\.\d+)?", str(value or ""))
        return max(0.0, float(match.group(0))) if match else 0.0

    def _apply_fill_opacity(self, shape: Any, opacity: Any) -> None:
        """CSS opacity (0..1) as transparency of a solid fill."""
        try:
            value = float(opacity)
        except (TypeError, ValueError):
            return
        if not 0 <= value < 1:
            return
        for color in shape.fill._xPr.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"):
            alpha = OxmlElement("a:alpha")
            alpha.set("val", str(int(round(value * 100000))))
            color.append(alpha)

    def _apply_shadow(self, shape: Any, box_shadow: Any) -> None:
        """No shadow unless the element has one, then a soft one like the CSS box-shadow.

        Shapes otherwise inherit the theme's shadow, which put a hard dark edge on every card and line."""
        shape.shadow.inherit = False
        # The shape's style also points at the theme's effect (a hard shadow), which LibreOffice still applies
        # when the effect list is empty: point it at "no effect".
        for ref in shape._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef"):
            ref.set("idx", "0")
        match = re.search(
            r"(-?\d+(?:\.\d+)?)(?:px)?\s+(-?\d+(?:\.\d+)?)(?:px)?\s+(\d+(?:\.\d+)?)(?:px)?[^r#]*(rgba?\([^)]+\)|#[0-9a-fA-F]{3,8})",
            str(box_shadow or ""),
        )
        if not match:
            return
        offset_x, offset_y, blur = (float(match.group(i)) for i in (1, 2, 3))
        color = match.group(4)
        alpha = 0.25
        rgba = re.match(r"rgba\(([^)]+)\)", color)
        if rgba:
            parts = [part.strip() for part in rgba.group(1).split(",")]
            if len(parts) == 4:
                try:
                    alpha = float(parts[3])
                except ValueError:
                    pass
        rgb = self._parse_color(color, RGBColor(15, 23, 42))
        distance = (offset_x**2 + offset_y**2) ** 0.5
        direction = math.degrees(math.atan2(offset_y, offset_x)) % 360 if distance else 90
        sp_pr = shape._element.spPr
        effects = sp_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        if effects is None:
            effects = OxmlElement("a:effectLst")
            sp_pr.append(effects)
        shadow = OxmlElement("a:outerShdw")
        shadow.set("blurRad", str(int(self._px_to_inches_w(blur))))
        shadow.set("dist", str(int(self._px_to_inches_w(distance))))
        shadow.set("dir", str(int(direction * 60000)))
        shadow.set("algn", "ctr")
        shadow.set("rotWithShape", "0")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", str(rgb))
        a = OxmlElement("a:alpha")
        # CSS shadows are spread over a large blur and look lighter than the same alpha in PowerPoint.
        a.set("val", str(int(max(0.0, min(1.0, alpha * 1.6)) * 100000)))
        clr.append(a)
        shadow.append(clr)
        effects.append(shadow)

    def _apply_line(self, shape: Any, styles: dict[str, Any], default_color: Optional[RGBColor] = None):
        border = str(styles.get("border") or "").strip()
        color = styles.get("borderColor") or styles.get("stroke")
        width = styles.get("borderWidth") or styles.get("strokeWidth")
        if border:
            color_match = re.search(r"#[0-9a-fA-F]{3,8}|rgba?\([^)]+\)", border)
            if color_match:
                color = color_match.group(0)
            width_match = re.search(r"(-?\d+(?:\.\d+)?)px", border)
            if width_match:
                width = width_match.group(1)
        if not color and not width:
            shape.line.fill.background()
            return
        if self._is_transparent(color):
            shape.line.fill.background()
            return
        shape.line.color.rgb = self._parse_color(color, default_color or self.theme["accent"])
        shape.line.width = Pt(max(0.25, float(self._parse_px(width, 1))))

    def _get_bullet_char(self, style: str, level: int) -> str:
        """Get the appropriate bullet character for the given style and level."""
        # CRITICAL FIX #5: Bullet character mapping based on SlideForge theme
        bullet_chars = {
            "default": ["•", "◦", "▪"],
            "square": ["■", "□", "▪"],
            "diamond": ["◆", "◇", "◈"],
            "modern": ["→", "–", "◆"],
            "chevron": ["»", "›", "–"],
            "dash": ["–", "—", "·"],
            "checklist": ["✓", "✓", "✓"],
            "star": ["✦", "✧", "•"],
        }
        chars = bullet_chars.get(style, bullet_chars["default"])
        return chars[min(level, len(chars) - 1)]

    def _set_paragraph_bullet(self, paragraph: Any, char: str = "•"):
        try:
            p_pr = paragraph._p.get_or_add_pPr()
            for child in list(p_pr):
                if child.tag.endswith("}buNone") or child.tag.endswith("}buChar") or child.tag.endswith("}buAutoNum"):
                    p_pr.remove(child)

            # A hanging indent per level, so the text starts clear of its marker and wrapped lines align with it.
            level = int(getattr(paragraph, "level", 0) or 0)
            p_pr.set("marL", str(int(Inches(0.3) * (level + 1))))
            p_pr.set("indent", str(-int(Inches(0.3))))

            # Use native PowerPoint auto-numbering if char is special indicator
            if char == "__AUTO_BULLET__":
                # Let PowerPoint handle bullets using default formatting
                bullet = OxmlElement("a:buFont")
                bullet.set("typeface", "Arial")
                p_pr.append(bullet)
            elif char.startswith("__AUTO_NUMBER__"):
                # Native PowerPoint numbering, in the scheme named after the marker (1. i. A. ...).
                buAutoNum = OxmlElement("a:buAutoNum")
                buAutoNum.set("type", char[len("__AUTO_NUMBER__"):].lstrip(":") or "arabicPeriod")
                p_pr.append(buAutoNum)
            else:
                # Custom bullet character
                bullet = OxmlElement("a:buChar")
                bullet.set("char", char)
                p_pr.append(bullet)
        except Exception:
            return

    def export(self) -> io.BytesIO:
        for index, slide_data in enumerate(self._normalized_slides()):
            try:
                self._add_slide(slide_data)
            except Exception:
                # One broken slide should not cost the user the whole deck.
                logger.exception("Skipping slide %s that could not be exported", index + 1)

        output = io.BytesIO()
        self.prs.save(output)
        output.seek(0)
        return output

    def _normalized_slides(self) -> list[Dict[str, Any]]:
        """Coerce loosely-typed saved state (null styles, string backgrounds, ...) into the shapes the exporter reads."""
        slides = self.state.get("slides") if isinstance(self.state, dict) else None
        normalized = []
        for slide in slides if isinstance(slides, list) else []:
            if not isinstance(slide, dict):
                continue
            slide = dict(slide)
            if not isinstance(slide.get("background"), dict):
                slide["background"] = None
            elements = []
            for element in slide.get("elements") if isinstance(slide.get("elements"), list) else []:
                if not isinstance(element, dict):
                    continue
                element = dict(element)
                if not isinstance(element.get("styles"), dict):
                    element["styles"] = {}
                elements.append(element)
            slide["elements"] = elements
            normalized.append(slide)
        return normalized

    @staticmethod
    def _z_index(element: Dict[str, Any]) -> float:
        styles = element.get("styles") if isinstance(element.get("styles"), dict) else {}
        try:
            return float(styles.get("zIndex", 0) or 0)
        except (TypeError, ValueError):
            return 0.0

    def _add_slide(self, slide_data: Dict[str, Any]):
        self.slide_bg_rgb = self.theme["background"]
        # Use a blank layout (index 6 is usually blank)
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # 1. Background
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.slide_bg_rgb
        gradient = self._theme_gradient()
        if gradient:
            try:
                fill.gradient()
                fill.gradient_angle = gradient[2]
                stops = fill.gradient_stops
                stops[0].color.rgb, stops[0].position = gradient[0], 0.0
                stops[-1].color.rgb, stops[-1].position = gradient[1], 1.0
            except Exception:
                fill.solid()
                fill.fore_color.rgb = self.slide_bg_rgb
        bg = slide_data.get("background")
        if bg and bg.get("content"):
            if bg.get("type") == "color":
                fill.solid()
                fill.fore_color.rgb = self._parse_color(bg["content"], self.slide_bg_rgb)
                self.slide_bg_rgb = fill.fore_color.rgb
            elif bg.get("type") == "image":
                img_stream = self._get_image_stream(bg["content"])
                if img_stream:
                    slide.shapes.add_picture(img_stream, 0, 0, width=self.prs.slide_width, height=self.prs.slide_height)
        elif slide_data.get("backgroundColor"):
            fill.solid()
            fill.fore_color.rgb = self._parse_color(slide_data.get("backgroundColor"), self.slide_bg_rgb)
            self.slide_bg_rgb = fill.fore_color.rgb

        if slide_data.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(slide_data.get("notes") or "")

        # 2. Elements
        elements = [*self._build_master_elements(slide_data), *slide_data.get("elements", [])]
        # Sort by zIndex
        sorted_elements = sorted(elements, key=self._z_index)

        shapes_of: dict[int, list] = {}
        for el in sorted_elements:
            # A preset's full-slide background box is drawn with the theme's slide background, which the slide
            # itself already has (gradient included); exported as its flat base colour it hid the gradient.
            if isinstance(el, dict) and el.get("presetBackground"):
                continue
            before = len(slide.shapes)
            self._add_element(slide, el)
            if isinstance(el, dict):
                shapes_of[id(el)] = list(slide.shapes)[before:]
                self._apply_rotation(shapes_of[id(el)], el)

        # Animations: PowerPoint export wrote none, so every object was simply there from the start.
        try:
            apply_slide_animations(slide, [el for el in slide_data.get("elements", []) if isinstance(el, dict)], shapes_of)
        except Exception:
            logger.exception("Could not export the slide's animations")

    def _apply_rotation(self, shapes: list, el: Dict[str, Any]) -> None:
        """The object's turn (the Rotation field). It was never written, so every shape came out level."""
        try:
            turn = float(el.get("rotation") or 0) % 360
        except (TypeError, ValueError):
            return
        if not turn or el.get("type") == "connector":
            return
        for shape in shapes:
            try:
                shape.rotation = turn
            except Exception:  # pictures of some kinds, graphic frames (tables, charts) cannot turn in python-pptx
                logger.debug("Could not rotate %s", getattr(shape, "shape_type", shape))

    def _master_config(self, master_id: str) -> Optional[dict[str, Any]]:
        if master_id == "none":
            return None
        defaults = {
            "content": {
                "id": "content",
                "enabled": True,
                "footerText": "Presentation",
                "logoText": "SlideForge",
                "showSlideNumber": True,
                "showFooter": True,
                "showTopRule": True,
            },
            "title": {
                "id": "title",
                "enabled": True,
                "footerText": "",
                "logoText": "SlideForge",
                "showSlideNumber": False,
                "showFooter": False,
                "showTopRule": True,
            },
            "section": {
                "id": "section",
                "enabled": True,
                "footerText": "Section",
                "logoText": "SlideForge",
                "showSlideNumber": True,
                "showFooter": True,
                "showTopRule": False,
            },
        }
        master_id = master_id if master_id in defaults else "content"
        masters = self.state.get("masterSlides") if isinstance(self.state.get("masterSlides"), dict) else {}
        config = {**defaults[master_id], **(masters.get(master_id) if isinstance(masters.get(master_id), dict) else {})}
        return config if config.get("enabled", True) else None

    def _master_shape(self, x: float, y: float, w: float, h: float, color: str, *, opacity: Optional[str] = None, border: Optional[str] = None, radius: str = "0px") -> dict[str, Any]:
        return {
            "type": "shape",
            "shapeType": "rectangle",
            "x": x,
            "y": y,
            "width": f"{w}px",
            "height": f"{h}px",
            "styles": {
                "backgroundColor": color,
                "borderRadius": radius,
                "zIndex": -20,
                **({"opacity": opacity} if opacity is not None else {}),
                **({"border": border} if border else {}),
            },
        }

    def _master_text(self, x: float, y: float, w: float, text: str, styles: dict[str, Any]) -> dict[str, Any]:
        return {
            "type": "text",
            "x": x,
            "y": y,
            "width": f"{w}px",
            "height": "28px",
            "content": text,
            "styles": {"zIndex": -19, "backgroundColor": "transparent", **styles},
        }

    def _build_master_elements(self, slide_data: dict[str, Any]) -> list[dict[str, Any]]:
        # The editor sends the layout decorations it draws as elements; drawing ours too would duplicate them.
        if self.state.get("masterElementsIncluded"):
            return []
        own_footer = bool(slide_data.get("editableMasterFooter")) or any(
            isinstance(el, dict) and el.get("editableMasterFooterElement") for el in slide_data.get("elements") or []
        )
        master_id = str(slide_data.get("masterId") or "content")
        config = self._master_config(master_id)
        if not config:
            return []
        try:
            slide_index = self.state.get("slides", []).index(slide_data)
        except ValueError:
            slide_index = 0
        slide_number = str(slide_index + 1).zfill(2)
        sx = self.base_w / 1024
        sy = self.base_h / 768
        text = self._rgb_to_hex(self.theme["text"])
        muted = self._rgb_to_hex(self.theme["muted"])
        accent = self._rgb_to_hex(self.theme["accent"])
        accent2 = self._rgb_to_hex(self.theme["accent2"])
        border = self._rgb_to_rgba(self.theme["muted"], 0.28)
        surface = self._rgb_to_rgba(self.theme["text"], 0.06)
        heading_font = "Aptos Display"
        body_font = DesignSystem.DEFAULT_FONT

        def sc(el: dict[str, Any]) -> dict[str, Any]:
            next_el = {**el, "x": round(float(el.get("x", 0)) * sx), "y": round(float(el.get("y", 0)) * sy)}
            next_el["width"] = f"{max(1, self._parse_px(el.get('width'), 1) * sx)}px"
            next_el["height"] = f"{max(1, self._parse_px(el.get('height'), 1) * sy)}px"
            styles = {**next_el.get("styles", {})}
            if styles.get("fontSize"):
                styles["fontSize"] = f"{max(1, self._parse_px(styles.get('fontSize'), 1) * min(sx, sy))}px"
            next_el["styles"] = styles
            return next_el

        elements: list[dict[str, Any]] = []
        if master_id == "title":
            if config.get("showTopRule", True):
                elements.append(self._master_shape(0, 0, 1024, 8, accent))
            elements.append(self._master_shape(882, 48, 88, 88, accent2, opacity="0.18", radius="22px"))
            elements.append(self._master_shape(930, 96, 40, 40, accent, opacity="0.82", radius="999px"))
            if config.get("logoText"):
                elements.append(self._master_text(64, 702, 420, str(config.get("logoText")), {
                    "color": muted, "fontFamily": body_font, "fontSize": "12px", "fontWeight": "700",
                }))
        elif master_id == "section":
            elements.append(self._master_shape(0, 0, 10, 768, accent))
            elements.append(self._master_shape(24, 42, 92, 92, surface, border=f"1px solid {border}", radius="24px"))
            elements.append(self._master_text(47, 67, 60, slide_number, {
                "color": accent, "fontFamily": heading_font, "fontSize": "34px", "fontWeight": "800", "textAlign": "center",
            }))
        else:
            if config.get("showTopRule", True):
                elements.append(self._master_shape(0, 0, 1024, 5, accent))
            elements.append({**self._master_shape(52, 712, 920, 1, border), "_footer": True})

        if config.get("showFooter", True):
            if config.get("logoText"):
                elements.append({**self._master_text(54, 724, 145, str(config.get("logoText")), {
                    "color": text, "fontFamily": body_font, "fontSize": "12px", "fontWeight": "800",
                }), "_footer": True})
            if config.get("footerText"):
                elements.append({**self._master_text(224, 724, 520, str(config.get("footerText")), {
                    "color": muted, "fontFamily": body_font, "fontSize": "12px", "fontWeight": "600",
                }), "_footer": True})
        if config.get("showSlideNumber", True):
            elements.append({**self._master_text(918, 724, 54, slide_number, {
                "color": muted, "fontFamily": body_font, "fontSize": "12px", "fontWeight": "800", "textAlign": "right",
            }), "_footer": True})
        # A slide whose footer is its own editable content keeps that footer instead of this one.
        return [sc(el) for el in elements if not (own_footer and el.get("_footer"))]

    def _rgb_to_hex(self, color: RGBColor) -> str:
        return f"#{color[0]:02X}{color[1]:02X}{color[2]:02X}"

    def _rgb_to_rgba(self, color: RGBColor, alpha: float) -> str:
        return f"rgba({color[0]},{color[1]},{color[2]},{alpha})"

    def _add_element(self, slide: Any, el: Dict[str, Any]):
        try:
            x = self._px_to_inches_w(float(el.get("x", 0)))
            y = self._px_to_inches_h(float(el.get("y", 0)))

            # Width and height can be "auto" or strings like "400px"
            w_raw = el.get("width", "100")
            h_raw = el.get("height", "100")

            w = self._px_to_inches_w(float(re.sub(r"[^\d.]", "", str(w_raw)) or 100))
            h = self._px_to_inches_h(float(re.sub(r"[^\d.]", "", str(h_raw)) or 100))

            el_type = el.get("type")
            if el_type == "text":
                self._add_text_element(slide, el, x, y, w, h)
            elif el_type == "image":
                self._add_image_element(slide, el, x, y, w, h)
            elif el_type == "shape":
                self._add_shape_element(slide, el, x, y, w, h)
            elif el_type == "table":
                self._add_table_element(slide, el, x, y, w, h)
            elif el_type == "connector":
                self._add_connector_element(slide, el, x, y, w, h)
            elif el_type == "video":
                self._add_video_element(slide, el, x, y, w, h)
            elif el_type == "mermaid":
                self._add_mermaid_element(slide, el, x, y, w, h)
            elif el_type == "chart":
                self._add_chart_element(slide, el, x, y, w, h)
            elif el_type in {"html", "pdf", "molecule", "equation", "latex"}:
                self._add_placeholder_element(slide, el, x, y, w, h)
            else:
                self._add_placeholder_element(slide, el, x, y, w, h)
        except Exception:
            logger.exception("Could not export element %s", el.get("id"))

    # CSS list-style-type -> DrawingML auto-number scheme (PowerPoint has no leading-zero or Greek numbering).
    NUMBER_SCHEMES = {
        "decimal": "arabicPeriod",
        "decimal-leading-zero": "arabicPeriod",
        "lower-roman": "romanLcPeriod",
        "upper-roman": "romanUcPeriod",
        "lower-alpha": "alphaLcPeriod",
        "lower-latin": "alphaLcPeriod",
        "upper-alpha": "alphaUcPeriod",
        "upper-latin": "alphaUcPeriod",
    }

    @staticmethod
    def _line_spacing(el: Dict[str, Any]):
        """Line spacing as a multiple of PowerPoint's single spacing (about 1.2 x the font size), from the line
        height the editor drew. None when the editor did not say, or the difference is not worth writing."""
        try:
            drawn = float(el.get("exportLineHeight") or (el.get("styles") or {}).get("lineHeight") or 0)
        except (TypeError, ValueError):
            return None
        if not 0.8 <= drawn <= 4:
            return None
        spacing = round(drawn / 1.2, 2)
        return None if abs(spacing - 1) < 0.04 else spacing

    def _add_text_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        styles = el.get("styles", {})
        content = el.get("content", "")

        # Create textbox
        shape = slide.shapes.add_textbox(x, y, w, h)
        if styles.get("backgroundColor") and not self._is_transparent(styles.get("backgroundColor")):
            self._apply_fill(shape, styles.get("backgroundColor"), self.slide_bg_rgb)
        else:
            shape.fill.background()
        self._apply_line(shape, styles, self.theme["accent"])
        tf = shape.text_frame
        tf.word_wrap = True
        # Only boxes the editor itself shrinks to fit ("autofit") are left to the office program to fit. For the
        # rest, "shrink text on overflow" made LibreOffice and PowerPoint re-fit every box with their own font
        # metrics: headline numbers and lists came out a fraction of their size.
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE if el.get("textFitMode") == "autofit" else MSO_AUTO_SIZE.NONE
        # The editor's text boxes have almost no inner margin; PowerPoint's default (0.1 in) shifted every text.
        tf.margin_left = tf.margin_right = self._px_to_inches_w(4)
        tf.margin_top = tf.margin_bottom = self._px_to_inches_h(2)

        # Handle structured bullet points or raw HTML
        tf.clear()
        line_spacing = self._line_spacing(el)
        try:
            paragraph_gap = max(0.0, min(80.0, float(el.get("exportParagraphGap") or 0)))
        except (TypeError, ValueError):
            paragraph_gap = 0.0
        html_items = html_list_items(content) if isinstance(content, str) else None
        if html_items:
            content = html_items
        if isinstance(content, list):
            # CRITICAL FIX #5: Enhanced list support for PPTX export
            for idx, item in enumerate(content):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.level = max(0, min(8, int(item.get("level", 0) or 0)))
                p.alignment = self._paragraph_alignment(styles)
                if line_spacing:
                    p.line_spacing = line_spacing
                if idx and paragraph_gap:
                    p.space_before = self._px_to_pt(paragraph_gap)
                
                # Set bullet based on list kind and style
                kind = item.get("kind", "bullet")
                style = item.get("style") or el.get("bulletStyle") or "default"
                
                if kind == "numbered":
                    scheme = self.NUMBER_SCHEMES.get(str(item.get("numberStyle") or "decimal"), "arabicPeriod")
                    self._set_paragraph_bullet(p, f"__AUTO_NUMBER__:{scheme}")
                else:
                    # Use bullet character based on style
                    char = self._get_bullet_char(style, p.level)
                    self._set_paragraph_bullet(p, char)
                
                raw = item.get("html") if isinstance(item, dict) and item.get("html") is not None else item.get("text", "")
                self._append_runs_to_paragraph(p, html_to_text_runs(raw), styles)
        else:
            p = tf.paragraphs[0]
            p.alignment = self._paragraph_alignment(styles)
            if line_spacing:
                p.line_spacing = line_spacing
            self._append_runs_to_paragraph(p, html_to_text_runs(content), styles)

    def _append_runs_to_paragraph(self, paragraph: Any, runs: list[dict[str, Any]], styles: dict[str, Any]):
        current = paragraph
        for run_data in runs:
            text = run_data.get("text", "")
            parts = text.split("\n")
            for idx, part in enumerate(parts):
                if idx > 0:
                    current = paragraph._parent.add_paragraph()
                    current.alignment = paragraph.alignment
                    if paragraph.line_spacing:
                        current.line_spacing = paragraph.line_spacing
                if not part:
                    continue
                run = current.add_run()
                run.text = part
                self._set_text_common(run, styles, extra=run_data)

    # Picture formats PowerPoint embeds as they are; anything else (WebP, AVIF, ...) is converted to PNG.
    PPTX_PICTURE_FORMATS = {"PNG", "JPEG", "GIF", "BMP", "TIFF"}

    def _picture_stream(self, content: str, width_px: float, height_px: float) -> Optional[io.BytesIO]:
        stream = self._get_image_stream(content)
        if not stream:
            return None
        raw = stream.getvalue()
        if str(content).startswith("data:image/svg") or raw[:512].lstrip().startswith((b"<svg", b"<?xml")):
            svg = self._sanitize_svg(raw.decode("utf-8", "ignore"))
            return self._svg_to_png_stream(svg, int(max(1, width_px) * 2), int(max(1, height_px) * 2))
        try:
            from PIL import Image

            with Image.open(io.BytesIO(raw)) as image:
                if image.format in self.PPTX_PICTURE_FORMATS:
                    return io.BytesIO(raw)
                converted = io.BytesIO()
                image.convert("RGBA").save(converted, "PNG")
                converted.seek(0)
                return converted
        except Exception:
            logger.warning("Could not read picture data for export")
            return None

    def _add_image_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        content = el.get("content", "")
        img_stream = self._picture_stream(content, self._parse_px(el.get("width"), 400), self._parse_px(el.get("height"), 300))
        if not img_stream:
            return

        pic = slide.shapes.add_picture(img_stream, x, y, width=w, height=h)
        self._apply_image_crop_transform(pic, el.get("cropTransform"))

    # Chart.js chart type -> PowerPoint chart type (bar charts turn sideways with indexAxis "y").
    CHART_TYPES = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
        "radar": XL_CHART_TYPE.RADAR_MARKERS,
        "polararea": XL_CHART_TYPE.PIE,
    }
    LEGEND_POSITIONS = {
        "top": XL_LEGEND_POSITION.TOP,
        "bottom": XL_LEGEND_POSITION.BOTTOM,
        "left": XL_LEGEND_POSITION.LEFT,
        "right": XL_LEGEND_POSITION.RIGHT,
    }

    @staticmethod
    def _chart_number(value: Any) -> Optional[float]:
        if isinstance(value, dict):
            value = value.get("y", value.get("value"))
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
        return number if number == number else None  # NaN -> gap

    def _add_chart_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        """A native, editable PowerPoint chart with the element's data, colours and legend."""
        data = el.get("chartData") if isinstance(el.get("chartData"), dict) else {}
        labels = [str(label) for label in data.get("labels") or []]
        datasets = [d for d in data.get("datasets") or [] if isinstance(d, dict)]
        if not labels or not datasets:
            self._add_placeholder_element(slide, el, x, y, w, h)
            return
        options = el.get("chartOptions") if isinstance(el.get("chartOptions"), dict) else {}
        kind = str(el.get("chartType") or "bar").lower()
        chart_type = self.CHART_TYPES.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
        if kind == "bar" and options.get("indexAxis") == "y":
            chart_type = XL_CHART_TYPE.BAR_CLUSTERED
        single_series = chart_type in {XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT}
        if single_series:
            datasets = datasets[:1]

        chart_data = CategoryChartData()
        chart_data.categories = labels
        for index, dataset in enumerate(datasets):
            values = list(dataset.get("data") or [])[: len(labels)]
            values += [None] * (len(labels) - len(values))
            chart_data.add_series(str(dataset.get("label") or f"Series {index + 1}"), [self._chart_number(v) for v in values])

        # The editor draws the chart on its own card (background, padding); keep that look.
        styles = el.get("styles", {}) or {}
        pad = self._parse_px(styles.get("padding"), 0) if styles.get("padding") else 0
        card_bg = styles.get("backgroundColor")
        has_card = bool(card_bg) and not self._is_transparent(card_bg)
        radius_px = self._corner_radius_px(styles.get("borderRadius"))
        if has_card:
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE if self._parse_px(styles.get("borderRadius"), 0) else MSO_SHAPE.RECTANGLE, x, y, w, h
            )
            if radius_px:
                card.adjustments[0] = max(0.0, min(0.5, radius_px / max(1.0, min(self._parse_px(el.get("width"), 500), self._parse_px(el.get("height"), 350)))))
            self._apply_fill(card, styles.get("backgroundColor"), RGBColor(255, 255, 255))
            card.line.fill.background()
            self._apply_shadow(card, styles.get("boxShadow"))
        inset_x = self._px_to_inches_w(pad) if pad else 0
        inset_y = self._px_to_inches_h(pad) if pad else 0
        frame = slide.shapes.add_chart(chart_type, x + inset_x, y + inset_y, w - 2 * inset_x, h - 2 * inset_y, chart_data)
        chart = frame.chart
        # chartStyle holds what the user set in the Chart Style panel; chartOptions what older decks stored.
        style = el.get("chartStyle") if isinstance(el.get("chartStyle"), dict) else {}
        try:
            font_px = max(8.0, min(40.0, float(style.get("fontSize") or 0) or 14.67))
        except (TypeError, ValueError):
            font_px = 14.67
        chart.font.size = Pt(round(font_px * 0.75, 1))
        # Text that reads on the chart's own card: the theme's light text was invisible on a white card.
        card_colour = self._parse_color(card_bg, RGBColor(255, 255, 255)) if has_card else self.slide_bg_rgb
        card_luminance = (0.2126 * card_colour[0] + 0.7152 * card_colour[1] + 0.0722 * card_colour[2]) / 255
        ink = RGBColor(0x33, 0x41, 0x55) if card_luminance > 0.5 else RGBColor(0xE2, 0xE8, 0xF0)
        chart.font.color.rgb = ink

        # No automatic title: PowerPoint and LibreOffice title a one-series chart with its series name, which
        # repeated the legend ("RMSD (Å)" twice). A title set in the editor is kept.
        plugins = options.get("plugins") if isinstance(options.get("plugins"), dict) else {}
        title = plugins.get("title") if isinstance(plugins.get("title"), dict) else {}
        title_text = str(style.get("title") or "").strip() or (str(title.get("text") or "").strip() if title.get("display") else "")
        chart.has_title = bool(title_text)
        if title_text:
            chart.chart_title.text_frame.text = title_text
            run = chart.chart_title.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(round((font_px + 3) * 0.75, 1))
            run.font.bold = True
            run.font.color.rgb = ink

        legend = plugins.get("legend") if isinstance(plugins.get("legend"), dict) else {}
        legend_position = str(style.get("legend") or "") or (str(legend.get("position") or "top") if legend.get("display", True) is not False else "none")
        chart.has_legend = legend_position != "none"
        if chart.has_legend:
            chart.legend.position = self.LEGEND_POSITIONS.get(legend_position, XL_LEGEND_POSITION.TOP)
            chart.legend.include_in_layout = False

        if not single_series:
            for axis, axis_title in ((chart.category_axis, style.get("xTitle")), (chart.value_axis, style.get("yTitle"))):
                text = str(axis_title or "").strip()
                if text:
                    axis.has_title = True
                    axis.axis_title.text_frame.text = text
                    run = axis.axis_title.text_frame.paragraphs[0].runs[0]
                    run.font.size = Pt(round(font_px * 0.75, 1))
                    run.font.bold = True
                    run.font.color.rgb = ink
            chart.value_axis.has_major_gridlines = style.get("grid") is not False
            if style.get("grid") is not False:
                chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0x94, 0xA3, 0xB8) if card_luminance > 0.5 else RGBColor(0x47, 0x55, 0x69)

        def colours(value: Any) -> list[str]:
            return [str(v) for v in value] if isinstance(value, list) else ([str(value)] if value else [])

        for series, dataset in zip(chart.plots[0].series, datasets):
            fills = colours(dataset.get("backgroundColor")) or colours(dataset.get("borderColor"))
            if not fills:
                continue
            if single_series:
                for point_index in range(len(labels)):
                    point = series.points[point_index]
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = self._parse_color(fills[point_index % len(fills)], self.theme["accent"])
            elif chart_type == XL_CHART_TYPE.LINE_MARKERS:
                line_colour = self._parse_color((colours(dataset.get("borderColor")) or fills)[0], self.theme["accent"])
                series.format.line.color.rgb = line_colour
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = line_colour
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = self._parse_color(fills[0], self.theme["accent"])
                if len(fills) > 1:  # one colour per bar, as Chart.js draws a colour array
                    for point_index in range(len(labels)):
                        point = series.points[point_index]
                        point.format.fill.solid()
                        point.format.fill.fore_color.rgb = self._parse_color(fills[point_index % len(fills)], self.theme["accent"])

    def _apply_image_crop_transform(self, pic: Any, crop: Optional[dict[str, Any]]):
        if not crop:
            return
        try:
            width_pct = max(100.0, float(crop.get("widthPercent", 100) or 100))
            height_pct = max(100.0, float(crop.get("heightPercent", 100) or 100))
            left_pct = min(0.0, max(100.0 - width_pct, float(crop.get("leftPercent", 0) or 0)))
            top_pct = min(0.0, max(100.0 - height_pct, float(crop.get("topPercent", 0) or 0)))
        except (TypeError, ValueError):
            return

        pic.crop_left = max(0.0, min(1.0, -left_pct / width_pct))
        pic.crop_right = max(0.0, min(1.0, (width_pct + left_pct - 100.0) / width_pct))
        pic.crop_top = max(0.0, min(1.0, -top_pct / height_pct))
        pic.crop_bottom = max(0.0, min(1.0, (height_pct + top_pct - 100.0) / height_pct))

    def _add_shape_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        shape_type_str = el.get("shapeType", "rectangle")

        mapping = {
            "rectangle": MSO_SHAPE.RECTANGLE,
            "circle": MSO_SHAPE.OVAL,
            "ellipse": MSO_SHAPE.OVAL,
            "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
            "diamond": MSO_SHAPE.DIAMOND,
            "hexagon": MSO_SHAPE.HEXAGON,
            "parallelogram": MSO_SHAPE.PARALLELOGRAM,
            "arrow-right": MSO_SHAPE.RIGHT_ARROW,
            "arrow-left": MSO_SHAPE.LEFT_ARROW,
            "arrow-up": MSO_SHAPE.UP_ARROW,
            "arrow-down": MSO_SHAPE.DOWN_ARROW,
            "right-triangle": MSO_SHAPE.RIGHT_TRIANGLE,
            "trapezoid": MSO_SHAPE.TRAPEZOID,
            "pentagon": MSO_SHAPE.REGULAR_PENTAGON,
            "octagon": MSO_SHAPE.OCTAGON,
            "star": MSO_SHAPE.STAR_5_POINT,
            "plus": MSO_SHAPE.CROSS,
            "chevron": MSO_SHAPE.CHEVRON,
            "callout": MSO_SHAPE.RECTANGULAR_CALLOUT,
        }

        mso_type = mapping.get(shape_type_str, MSO_SHAPE.RECTANGLE)
        styles = el.get("styles", {})
        width_px = self._parse_px(el.get("width"), 100)
        height_px = self._parse_px(el.get("height"), 100)
        radius_px = self._corner_radius_px(styles.get("borderRadius"))
        rounded = mso_type == MSO_SHAPE.RECTANGLE and radius_px > 0
        if rounded:
            mso_type = MSO_SHAPE.ROUNDED_RECTANGLE
        # Words in a non-rectangular shape go in a text box over the editor's text area, grouped with the shape.
        # PowerPoint fits a shape's own text into the preset's inner rectangle, which for a star or triangle is
        # far smaller than the editor's, so short words broke mid-word ("Nati / ve").
        has_text = bool(str(el.get("shapeText") or "").strip())
        overlay = has_text and mso_type not in {MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE}
        container = slide.shapes.add_group_shape().shapes if overlay else slide.shapes
        shape = container.add_shape(mso_type, x, y, w, h)
        if rounded:
            # The corner radius as a share of the shorter side (0.5 = fully round: pills and circles).
            shape.adjustments[0] = max(0.0, min(0.5, radius_px / max(1.0, min(width_px, height_px))))

        self._apply_fill(shape, styles.get("backgroundColor"), self.theme["accent"])
        self._apply_fill_opacity(shape, styles.get("opacity"))
        self._apply_line(shape, styles, self.theme["accent"])
        self._apply_shadow(shape, styles.get("boxShadow"))
        if overlay:
            top, right, bottom, left = self.SHAPE_TEXT_INSETS.get(shape_type_str, (8, 10, 8, 10))
            box = container.add_textbox(
                x + int(w * left / 100), y + int(h * top / 100), int(w * (100 - left - right) / 100), int(h * (100 - top - bottom) / 100)
            )
            self._add_shape_text(box, el, margin=Inches(0))
        else:
            self._add_shape_text(shape, el)

    # The editor's text area in each shape, in percent of the box (top, right, bottom, left): shapes.js _shapeTextInset.
    SHAPE_TEXT_INSETS = {
        "triangle": (42, 18, 6, 18),
        "right-triangle": (40, 40, 6, 6),
        "diamond": (22, 22, 22, 22),
        "star": (30, 26, 22, 26),
        "callout": (6, 8, 30, 8),
        "pentagon": (18, 14, 8, 14),
        "chevron": (8, 26, 8, 26),
        "parallelogram": (8, 22, 8, 22),
        "trapezoid": (8, 22, 8, 22),
    }

    def _add_shape_text(self, shape: Any, el: Dict[str, Any], margin: Any = None):
        """Words typed into a shape go into the shape's own text frame (or the text box laid over it)."""
        text = str(el.get("shapeText") or "")
        if not text.strip():
            return
        style = el.get("shapeTextStyle") if isinstance(el.get("shapeTextStyle"), dict) else {}
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}.get(str(style.get("verticalAlign")), MSO_ANCHOR.MIDDLE)
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(frame, side, Inches(0.06) if margin is None else margin)
        align = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}.get(str(style.get("textAlign")), PP_ALIGN.CENTER)
        size = Pt(max(6.0, self._parse_px(style.get("fontSize"), 20) * 0.75))
        colour = self._parse_color(style.get("color"), RGBColor(0xFF, 0xFF, 0xFF))
        bold = str(style.get("fontWeight") or "") in {"600", "700", "800", "900", "bold"}
        italic = str(style.get("fontStyle") or "") == "italic"
        font_family = str(style.get("fontFamily") or "").split(",")[0].replace('"', "").replace("'", "").strip() or None
        for index, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            run = paragraph.add_run()
            run.text = line
            run.font.size = size
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = colour
            if font_family:
                run.font.name = font_family

    def _add_table_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        table_data = el.get("tableData") or {}
        rows = max(1, int(table_data.get("rows") or len(table_data.get("cells") or []) or 1))
        cols = max(1, int(table_data.get("cols") or max((len(row) for row in table_data.get("cells", []) if isinstance(row, list)), default=1)))
        rows = min(rows, 50)
        cols = min(cols, 20)
        table_shape = slide.shapes.add_table(rows, cols, x, y, w, h)
        table = table_shape.table
        cells = table_data.get("cells") if isinstance(table_data.get("cells"), list) else []
        header_row = table_data.get("headerRow", True)
        font_size = table_data.get("fontSize") or "14px"

        # In the editor the grid stretches to the element's box; its stored column and row sizes are proportions.
        col_widths = [max(1.0, float(w_ or 140)) for w_ in (table_data.get("colWidths") or [])][:cols]
        col_widths += [140.0] * (cols - len(col_widths))
        row_heights = [max(1.0, float(h_ or 44)) for h_ in (table_data.get("rowHeights") or [])][:rows]
        row_heights += [44.0] * (rows - len(row_heights))
        for col_idx, width in enumerate(col_widths):
            table.columns[col_idx].width = int(w * width / sum(col_widths))
        for row_idx, height in enumerate(row_heights):
            table.rows[row_idx].height = int(h * height / sum(row_heights))
        border_colour = str(self._parse_color(table_data.get("borderColor") or "#cbd5e1", RGBColor(0xCB, 0xD5, 0xE1)))
        border_width = int(Pt(max(0.25, float(self._parse_px(table_data.get("borderWidth"), 1)) * 0.75)))

        for row_idx in range(rows):
            for col_idx in range(cols):
                cell = table.cell(row_idx, col_idx)
                raw_cell = cells[row_idx][col_idx] if row_idx < len(cells) and isinstance(cells[row_idx], list) and col_idx < len(cells[row_idx]) else {}
                text = raw_cell.get("text", "") if isinstance(raw_cell, dict) else str(raw_cell or "")
                cell_styles = raw_cell.get("styles", {}) if isinstance(raw_cell, dict) and isinstance(raw_cell.get("styles"), dict) else {}
                cell.text = str(text)
                tc_pr = cell._tc.get_or_add_tcPr()
                for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                    line = OxmlElement(side)
                    line.set("w", str(border_width))
                    solid = OxmlElement("a:solidFill")
                    colour = OxmlElement("a:srgbClr")
                    colour.set("val", border_colour)
                    solid.append(colour)
                    line.append(solid)
                    tc_pr.append(line)
                if row_idx == 0 and header_row:
                    fill_color = cell_styles.get("backgroundColor") or table_data.get("headerFill") or "#e2e8f0"
                elif table_data.get("zebra") and row_idx % 2 == 0:
                    fill_color = cell_styles.get("backgroundColor") or table_data.get("altFill") or "#f8fafc"
                else:
                    fill_color = cell_styles.get("backgroundColor") or table_data.get("bodyFill") or "#ffffff"
                cell.fill.solid()
                cell.fill.fore_color.rgb = self._parse_color(fill_color, self.slide_bg_rgb)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = self._paragraph_alignment({"textAlign": cell_styles.get("textAlign") or table_data.get("textAlign", "left")})
                    for run in paragraph.runs:
                        self._set_text_common(
                            run,
                            {
                                "fontSize": cell_styles.get("fontSize") or font_size,
                                "fontFamily": cell_styles.get("fontFamily") or table_data.get("fontFamily", DesignSystem.DEFAULT_FONT),
                                "fontWeight": cell_styles.get("fontWeight") or ("700" if row_idx == 0 and header_row else table_data.get("fontWeight", "400")),
                                "fontStyle": cell_styles.get("fontStyle") or table_data.get("fontStyle"),
                                "color": cell_styles.get("color") or table_data.get("headerTextColor" if row_idx == 0 and header_row else "textColor") or ("#ffffff" if row_idx == 0 and header_row else None),
                            },
                        )
    # Editor arrowheads -> DrawingML line ends (PowerPoint has no square end; a dot is the nearest).
    CONNECTOR_HEADS = {"arrow": "triangle", "triangle": "triangle", "chevron": "arrow", "dot": "oval", "square": "oval", "diamond": "diamond"}

    @staticmethod
    def _connector_path_points(points: list[tuple[float, float]], kind: str) -> list[tuple[float, float]]:
        """The connector's path as a polyline, as the editor draws it (a curve is quadratic segments through midpoints)."""
        if kind != "curve" or len(points) < 3:
            return points if kind == "poly" else [points[0], points[-1]]

        def quad(a, c, b, steps=16):
            return [((1 - t) ** 2 * a[0] + 2 * (1 - t) * t * c[0] + t ** 2 * b[0], (1 - t) ** 2 * a[1] + 2 * (1 - t) * t * c[1] + t ** 2 * b[1]) for t in (i / steps for i in range(1, steps + 1))]

        out = [points[0]]
        current, control = points[0], points[0]
        for i in range(1, len(points) - 1):
            mid = ((points[i][0] + points[i + 1][0]) / 2, (points[i][1] + points[i + 1][1]) / 2)
            out += quad(current, points[i], mid)
            current, control = mid, points[i]
        reflected = (2 * current[0] - control[0], 2 * current[1] - control[1])  # SVG "T": the last control mirrored
        out += quad(current, reflected, points[-1])
        return out

    def _add_connector_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        """The connector along its own points, with its arrowheads. It was a diagonal of its box, which missed the
        shapes it joins (the box has padding around the line) and pointed the wrong way for a rising line."""
        styles = el.get("styles", {}) or {}
        width_px = max(1.0, self._parse_px(el.get("width"), 280))
        height_px = max(1.0, self._parse_px(el.get("height"), 140))
        raw = [p for p in (el.get("points") or []) if isinstance(p, dict)]
        points: list[tuple[float, float]] = []
        for p in raw:
            try:
                points.append((float(p.get("x")), float(p.get("y"))))
            except (TypeError, ValueError):
                continue
        if len(points) < 2:
            points = [(0.0, 0.0), (width_px, height_px)]
        kind = str(el.get("connectorType") or "line")
        path = self._connector_path_points(points, kind)
        emu = [(int(x + w * (px / width_px)), int(y + h * (py / height_px))) for px, py in path]
        colour = self._parse_color(styles.get("color") or styles.get("borderColor"), self.theme["accent"])
        stroke = Pt(float(self._parse_px(styles.get("strokeWidth") or styles.get("borderWidth"), 2)))
        if len(emu) == 2:
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu[0][0], emu[0][1], emu[1][0], emu[1][1])
        else:
            builder = slide.shapes.build_freeform(emu[0][0], emu[0][1], scale=1.0)
            builder.add_line_segments(emu[1:], close=False)
            shape = builder.convert_to_shape()
            shape.fill.background()
            shape.shadow.inherit = False
        shape.line.color.rgb = colour
        shape.line.width = stroke
        ln = shape.line._get_or_add_ln()
        for tag, head in (("a:headEnd", el.get("connectorStart") or "none"), ("a:tailEnd", el.get("connectorEnd") or "arrow")):
            kind_name = self.CONNECTOR_HEADS.get(str(head))
            if not kind_name:
                continue
            end = OxmlElement(tag)
            end.set("type", kind_name)
            end.set("w", "med")
            end.set("len", "med")
            ln.append(end)

    def _sanitize_svg(self, svg: str) -> str:
        cleaned = str(svg or "")
        cleaned = re.sub(r"<\s*(script|foreignObject)\b[^>]*>.*?<\s*/\s*\1\s*>", "", cleaned, flags=re.I | re.S)
        cleaned = re.sub(r"\s+on[a-zA-Z]+\s*=\s*(['\"]).*?\1", "", cleaned, flags=re.I | re.S)
        cleaned = re.sub(r"\s+(href|xlink:href)\s*=\s*(['\"])\s*(javascript:|data:text/html).*?\2", "", cleaned, flags=re.I | re.S)
        cleaned = re.sub(r"url\s*\(\s*(['\"]?)\s*(javascript:|data:text/html).*?\)", "none", cleaned, flags=re.I | re.S)
        return cleaned.strip()

    def _svg_to_png_stream(self, svg: str, width_px: int, height_px: int) -> Optional[io.BytesIO]:
        try:
            import cairosvg

            png = cairosvg.svg2png(
                bytestring=svg.encode("utf-8"),
                output_width=max(1, width_px),
                output_height=max(1, height_px),
            )
            stream = io.BytesIO(png)
            stream.seek(0)
            return stream
        except Exception:
            pass

        try:
            from svglib.svglib import svg2rlg
            from reportlab.graphics import renderPM

            drawing = svg2rlg(io.BytesIO(svg.encode("utf-8")))
            png = renderPM.drawToString(drawing, fmt="PNG", dpi=220)
            stream = io.BytesIO(png)
            stream.seek(0)
            return stream
        except Exception:
            return None

    def _add_mermaid_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        svg = self._sanitize_svg(el.get("svgContent") or "")
        if not svg:
            self._add_placeholder_element(slide, {**el, "type": "mermaid"}, x, y, w, h)
            return

        svg_stream = io.BytesIO(svg.encode("utf-8"))
        try:
            slide.shapes.add_picture(svg_stream, x, y, width=w, height=h)
            return
        except Exception:
            pass

        width_px = int(max(320, self._parse_px(el.get("width"), 560) * 3))
        height_px = int(max(240, self._parse_px(el.get("height"), 360) * 3))
        png_stream = self._svg_to_png_stream(svg, width_px, height_px)
        if png_stream:
            slide.shapes.add_picture(png_stream, x, y, width=w, height=h)
            return
        self._add_placeholder_element(slide, {**el, "type": "mermaid"}, x, y, w, h)

    def _add_placeholder_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        label_by_type = {
            "video": "Video",
            "html": "HTML embed",
            "pdf": "PDF",
            "molecule": "3D molecule",
            "chart": "Chart",
            "equation": "Equation",
            "latex": "Equation",
            "mermaid": "Mermaid diagram",
        }
        label = label_by_type.get(el.get("type"), "Unsupported content")
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._parse_color("rgba(148, 163, 184, 0.18)", self.theme["muted"])
        shape.line.color.rgb = self.theme["muted"]
        tf = shape.text_frame
        tf.text = f"{label} placeholder"
        paragraph = tf.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        if paragraph.runs:
            run = paragraph.runs[0]
        else:
            run = paragraph.add_run()
            run.text = f"{label} placeholder"
        run.font.size = Pt(14)
        run.font.color.rgb = self.theme["muted"]

    def _resolve_local_video_path(self, content: str) -> Optional[Path]:
        normalized = str(content or "").strip()
        if not normalized or "://" in normalized:
            return None

        candidates: list[Path] = []
        if normalized.startswith("/media/"):
            candidates.extend(self._media_candidates(normalized))
        elif normalized.startswith("/static/"):
            candidates.append(self.project_root / normalized.lstrip("/"))
        elif normalized.startswith("/"):
            candidates.append(Path(normalized))
        else:
            candidates.append(self.project_root / normalized)

        safe_roots = self._safe_asset_roots()
        for candidate in candidates:
            try:
                path = candidate.expanduser().resolve()
            except Exception:
                continue
            if path.exists() and path.is_file():
                if any(path == root or path.is_relative_to(root) for root in safe_roots):
                    return path
        return None

    def _add_video_element(self, slide: Any, el: Dict[str, Any], x: Inches, y: Inches, w: Inches, h: Inches):
        content = el.get("content", "")
        video_path = self._resolve_local_video_path(content)
        if video_path:
            try:
                slide.shapes.add_movie(
                    str(video_path),
                    x,
                    y,
                    width=w,
                    height=h,
                    mime_type="video/mp4"
                )
                return
            except Exception as e:
                logger.warning("Could not embed video file in PPTX: %s", e)

        # Fall back to visual placeholder if file resolution or embedding fails
        self._add_placeholder_element(slide, el, x, y, w, h)
