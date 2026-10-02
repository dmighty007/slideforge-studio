import os
from django.test import TestCase
from unittest import mock

from slideforge.bridge.llm_utils import _choose_text_model
from slideforge.bridge.pdf_bridge import PDF2PPTxBridge

class LLMUtilsTests(TestCase):
    @mock.patch("slideforge.bridge.llm_utils._available_ollama_models")
    def test_choose_text_model_with_env(self, mock_available):
        with mock.patch.dict(os.environ, {"PPTMAKER_TEXT_MODEL": "custom-model:latest"}):
            self.assertEqual(_choose_text_model(), "custom-model:latest")
            
    @mock.patch("slideforge.bridge.llm_utils._available_ollama_models")
    def test_choose_text_model_fallback(self, mock_available):
        mock_available.return_value = {"llama3.1:8b"}
        with mock.patch.dict(os.environ, clear=True):
            self.assertEqual(_choose_text_model(), "llama3.1:8b")

class PDFTextCleaningTests(TestCase):
    def test_text_from_pdfs_keeps_greek_accents_and_symbols_and_expands_ligatures(self):
        from slideforge.bridge.pdf_text_extractor import _clean_text

        # It was transliterated to ASCII: "a-helix", "Muller", "A", and no maths symbols.
        self.assertEqual(_clean_text("α-helix  of\u00adMüller,  5 Å, ∑ x² ≤ 1"), "α-helix ofMüller, 5 Å, ∑ x² ≤ 1")
        self.assertEqual(_clean_text("ﬁgure ﬂow ﬃ"), "figure flow ffi")
        self.assertEqual(_clean_text("蛋白质结构"), "蛋白质结构")


class PDFBridgeTests(TestCase):
    def test_clean_generated_text(self):
        # Should strip prompt remnants
        self.assertEqual(PDF2PPTxBridge._clean_generated_text("Explain the process"), "")
        # Should strip generic visual text
        self.assertEqual(PDF2PPTxBridge._clean_generated_text("This diagram illustrates the process"), "")
        # Normal text should pass through
        self.assertEqual(PDF2PPTxBridge._clean_generated_text("The results show an increase in accuracy."), "The results show an increase in accuracy.")
        # Should strip untitled presentation
        self.assertEqual(PDF2PPTxBridge._clean_generated_text("Untitled Presentation"), "")


class PPTXExportTextTests(TestCase):
    def _export(self, state):
        from pptx import Presentation
        from slideforge.bridge.pptx_exporter import PPTXExporter

        return Presentation(PPTXExporter(state).export())

    def _texts(self, slide):
        return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]

    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

    def _click_steps(self, slide):
        timing = slide._element.find(f"{self.P}timing")
        if timing is None:
            return None, []
        main = next(node for node in timing.iter(f"{self.P}cTn") if node.get("nodeType") == "mainSeq")
        return timing, list(main.find(f"{self.P}childTnLst"))

    def test_rotation_is_written_to_the_shapes(self):
        prs = self._export({"pageSetup": "talk-16-9", "slides": [{"elements": [
            {"type": "shape", "shapeType": "rectangle", "x": 100, "y": 100, "width": "200px", "height": "100px", "rotation": 45, "styles": {"backgroundColor": "#2563eb"}},
            {"type": "text", "x": 500, "y": 100, "width": "300px", "height": "60px", "rotation": -30, "content": "Turned", "styles": {"fontSize": "24px"}},
        ]}]})
        turns = sorted(round(shape.rotation) for shape in prs.slides[0].shapes if shape.width > 1500000 and shape.height > 500000 and shape.width < 4000000)
        self.assertEqual(turns, [45, 330])

    def test_links_in_text_become_powerpoint_hyperlinks(self):
        prs = self._export({"pageSetup": "talk-16-9", "slides": [{"elements": [
            {"type": "text", "x": 100, "y": 100, "width": "500px", "height": "60px", "styles": {"fontSize": "24px"},
             "content": 'See <a href="https://example.org/paper">the paper</a> and <a href="javascript:alert(1)">this</a>'},
        ]}]})
        runs = [run for shape in prs.slides[0].shapes if shape.has_text_frame for p in shape.text_frame.paragraphs for run in p.runs]
        links = {run.text: run.hyperlink.address for run in runs if run.text.strip()}
        self.assertEqual(links.get("the paper"), "https://example.org/paper")
        self.assertIsNone(links.get("this"))
        self.assertIsNone(links.get("See"))

    def test_click_animations_become_powerpoint_click_effects(self):
        import io

        from pptx import Presentation

        fade = {"id": "a1", "category": "entrance", "type": "fadeIn", "duration": 600, "delay": 0, "startOpacity": 0, "endOpacity": 1, "trigger": "on-click"}
        fly = {**fade, "id": "a2", "type": "transform", "direction": "up"}
        auto = {**fade, "id": "a3", "trigger": "on-slide"}
        pulse = {"id": "a4", "category": "emphasis", "type": "scaleInPlace", "duration": 600, "startScale": 1, "endScale": 1.1, "trigger": "on-click"}
        box = lambda x, animations: {
            "type": "shape", "shapeType": "rectangle", "x": x, "y": 100, "width": "120px", "height": "80px", "styles": {"backgroundColor": "#2563eb"},
            "animation": {"timelines": [{"animations": animations}]},
        }
        prs = self._export({"pageSetup": "talk-16-9", "slides": [
            {"elements": [box(0, [fade]), box(200, [fly]), box(400, [auto]), box(600, [pulse])]},
            {"elements": [box(0, [pulse])]},
        ]})
        slide = prs.slides[0]
        timing, steps = self._click_steps(slide)
        self.assertEqual(len(steps), 3)  # the automatic step, then two clicks; the emphasis effect is left out
        boxes = sorted((shape for shape in slide.shapes if shape.shape_type == 1 and round(shape.width / 9525) == 120), key=lambda shape: shape.left)
        ids = [shape.shape_id for shape in boxes]
        effects = [node for node in timing.iter(f"{self.P}cTn") if node.get("presetClass")]
        self.assertEqual([(e.get("nodeType"), e.get("presetID"), e.get("presetSubtype")) for e in effects],
                         [("afterEffect", "10", "0"), ("clickEffect", "10", "0"), ("clickEffect", "2", "4")])
        targets = [node.get("spid") for node in timing.iter(f"{self.P}spTgt")]
        self.assertEqual(set(targets), {str(ids[0]), str(ids[1]), str(ids[2])})
        self.assertIsNone(prs.slides[1]._element.find(f"{self.P}timing"))
        # Still a file PowerPoint's own reader accepts, with the timing after the shapes.
        buffer = io.BytesIO()
        prs.save(buffer)
        again = Presentation(io.BytesIO(buffer.getvalue()))
        self.assertEqual([child.tag.split("}")[1] for child in again.slides[0]._element][-1], "timing")

    def test_bullet_reveal_becomes_a_paragraph_build(self):
        fade = {"id": "a1", "category": "entrance", "type": "fadeIn", "duration": 600, "delay": 0, "startOpacity": 0, "endOpacity": 1, "trigger": "on-click"}
        prs = self._export({"pageSetup": "talk-16-9", "slides": [{"elements": [{
            "type": "text", "x": 100, "y": 100, "width": "600px", "height": "200px", "revealBullets": True,
            "content": [{"text": "First", "level": 0}, {"text": "Detail", "level": 1}, {"text": "Second", "level": 0}],
            "styles": {"fontSize": "20px"}, "animation": {"timelines": [{"animations": [fade]}]},
        }]}]})
        slide = prs.slides[0]
        timing, steps = self._click_steps(slide)
        self.assertEqual(len(steps), 2)  # "First" with its detail, then "Second"; no extra click for the box
        ranges = [[(r.get("st"), r.get("end")) for r in step.iter(f"{self.P}pRg")] for step in steps]
        self.assertEqual([sorted(set(r)) for r in ranges], [[("0", "0"), ("1", "1")], [("2", "2")]])
        build = timing.find(f"{self.P}bldLst/{self.P}bldP")
        self.assertEqual(build.get("build"), "p")
        self.assertEqual({node.get("presetID") for node in timing.iter(f"{self.P}cTn") if node.get("presetClass")}, {"10"})

    def test_preset_background_box_leaves_the_theme_gradient_showing(self):
        from pptx.enum.dml import MSO_FILL

        prs = self._export({
            "pageSetup": "talk-16-9",
            "themeBackgroundCss": "linear-gradient(180deg, #08111F 0%, #1E3A5F 100%)",
            "slides": [{"elements": [
                {"type": "shape", "shapeType": "rectangle", "presetBackground": True, "x": 0, "y": 0, "width": "1280px", "height": "720px",
                 "styles": {"backgroundColor": "#0B1020", "background": "linear-gradient(180deg, #08111F 0%, #1E3A5F 100%)"}},
                {"type": "text", "x": 64, "y": 80, "width": "600px", "height": "60px", "content": "Title", "styles": {"fontSize": "34px"}},
            ]}],
        })
        slide = prs.slides[0]
        self.assertEqual(slide.background.fill.type, MSO_FILL.GRADIENT)
        full = [s for s in slide.shapes if s.width >= prs.slide_width * 0.95 and s.height >= prs.slide_height * 0.95]
        self.assertEqual(full, [], "a full-slide box covers the theme gradient")
        self.assertAlmostEqual(slide.background.fill.gradient_angle, 270.0)
        # A diagonal theme gradient (Graphite's 155deg) is kept too, at its angle.
        diagonal = self._export({
            "pageSetup": "talk-16-9",
            "themeBackgroundCss": "repeating-linear-gradient(90deg, rgba(255,255,255,0.035) 0 1px, transparent 1px 72px), linear-gradient(155deg, #090E1A 0%, #111827 52%, #202938 100%)",
            "slides": [{"elements": []}],
        }).slides[0].background.fill
        self.assertEqual(diagonal.type, MSO_FILL.GRADIENT)
        self.assertAlmostEqual(diagonal.gradient_angle, 295.0)
        self.assertEqual(str(diagonal.gradient_stops[0].color.rgb), "090E1A")

    def test_numbered_html_list_exports_native_numbering_with_levels(self):
        content = (
            '<ol class="ppt-numbered-block" style="list-style-type:lower-roman;"><li>First</li>'
            '<li>Second<ol style="list-style-type:lower-roman;"><li>Nested</li></ol></li></ol>'
        )
        state = {"slides": [{"masterId": "none", "elements": [{"type": "text", "content": content, "styles": {}}]}]}
        slide = self._export(state).slides[0]
        shape = next(s for s in slide.shapes if s.has_text_frame and "First" in s.text_frame.text)
        paragraphs = shape.text_frame.paragraphs
        self.assertEqual([p.text for p in paragraphs], ["First", "Second", "Nested"])
        self.assertEqual([p.level for p in paragraphs], [0, 0, 1])
        for paragraph in paragraphs:
            numbering = [c for c in paragraph._p.pPr if c.tag.endswith("}buAutoNum")]
            self.assertEqual(numbering[0].get("type"), "romanLcPeriod")

    def test_structured_bullets_use_the_element_bullet_style(self):
        state = {"slides": [{"masterId": "none", "elements": [
            {"type": "text", "bulletStyle": "diamond", "content": [{"html": "Point", "level": 0}], "styles": {}},
        ]}]}
        paragraph = next(s for s in self._export(state).slides[0].shapes if s.has_text_frame).text_frame.paragraphs[0]
        self.assertEqual([c.get("char") for c in paragraph._p.pPr if c.tag.endswith("}buChar")], ["◆"])

    def test_slide_with_editable_footer_is_not_given_a_second_footer(self):
        footer = {"type": "text", "content": "SlideForge", "editableMasterFooterElement": True, "styles": {}}
        state = {"slides": [{"masterId": "content", "editableMasterFooter": True, "elements": [footer]}]}
        self.assertEqual(self._texts(self._export(state).slides[0]).count("SlideForge"), 1)

    def test_editor_sent_master_elements_are_not_redrawn(self):
        state = {"masterElementsIncluded": True, "slides": [{"masterId": "content", "elements": []}]}
        self.assertEqual(self._texts(self._export(state).slides[0]), [])


class PPTXExportFidelityTests(TestCase):
    """What the export must keep for a deck to look in PowerPoint or LibreOffice as it does in the editor."""

    def _prs(self, elements, **state):
        from pptx import Presentation
        from slideforge.bridge.pptx_exporter import PPTXExporter

        return Presentation(PPTXExporter({"slides": [{"masterId": "none", "elements": elements}], **state}).export())

    def _slide(self, elements, **state):
        return self._prs(elements, **state).slides[0]

    def test_every_page_setup_of_the_editor_gives_its_own_slide_size(self):
        from pptx.util import Inches

        for setup, (width, height) in {"talk-16-9": (1280, 720), "lecture-16-10": (1280, 800), "standard-4-3": (1024, 768)}.items():
            prs = self._prs([], pageSetup=setup)
            self.assertEqual((prs.slide_width, prs.slide_height), (Inches(width / 96), Inches(height / 96)), setup)

    def test_theme_gradient_becomes_a_gradient_slide_background(self):
        css = "linear-gradient(180deg, #07111f 0%, #1a3a6b 100%)"
        fill = self._slide([], themeBackgroundCss=css).background.fill
        self.assertEqual([str(stop.color.rgb) for stop in fill.gradient_stops], ["07111F", "1A3A6B"])

    def test_text_keeps_its_size_and_list_spacing(self):
        from pptx.enum.text import MSO_AUTO_SIZE
        from pptx.util import Inches

        slide = self._slide([{
            "type": "text", "x": 0, "y": 0, "width": "400px", "height": "200px", "styles": {"fontSize": "24px"},
            "content": [{"html": "One", "level": 0}, {"html": "Two", "level": 1}],
            "exportLineHeight": 1.5, "exportParagraphGap": 8,
        }])
        frame = next(s for s in slide.shapes if s.has_text_frame).text_frame
        # "Shrink text on overflow" let the office program re-fit every box with its own font metrics.
        self.assertEqual(frame.auto_size, MSO_AUTO_SIZE.NONE)
        first, second = frame.paragraphs
        self.assertEqual(first.line_spacing, 1.25)
        self.assertEqual(second.space_before.pt, 6)
        # The bullet hangs in front of the text instead of touching it.
        self.assertEqual(int(first._p.pPr.get("marL")), Inches(0.3))
        self.assertEqual(int(first._p.pPr.get("indent")), -Inches(0.3))
        self.assertEqual(int(second._p.pPr.get("marL")), Inches(0.6))

    def test_card_keeps_round_corners_and_takes_no_shadow_from_the_office_theme(self):
        from pptx.enum.shapes import MSO_SHAPE

        def card(styles):
            slide = self._slide([{"type": "shape", "shapeType": "rectangle", "x": 0, "y": 0, "width": "400px", "height": "200px", "styles": styles}])
            return next(s for s in slide.shapes if s.shape_type == 1)

        rounded = card({"backgroundColor": "#ffffff", "borderRadius": "20px", "boxShadow": "0 8px 24px rgba(15, 23, 42, 0.12)"})
        self.assertEqual(rounded.auto_shape_type, MSO_SHAPE.ROUNDED_RECTANGLE)
        self.assertAlmostEqual(rounded.adjustments[0], 0.1, places=2)
        self.assertEqual(len(rounded._element.xpath(".//a:outerShdw")), 1)
        plain = card({"backgroundColor": "#ffffff"})
        self.assertEqual(plain.auto_shape_type, MSO_SHAPE.RECTANGLE)
        self.assertEqual(plain._element.xpath(".//a:outerShdw"), [])
        # LibreOffice draws the theme's shadow for any shape whose style still points at it.
        self.assertEqual([ref.get("idx") for ref in plain._element.xpath(".//a:effectRef")], ["0"])

    def test_every_editor_shape_becomes_the_matching_powerpoint_shape(self):
        from pptx.enum.shapes import MSO_SHAPE

        expected = {
            "circle": MSO_SHAPE.OVAL, "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "right-triangle": MSO_SHAPE.RIGHT_TRIANGLE,
            "diamond": MSO_SHAPE.DIAMOND, "parallelogram": MSO_SHAPE.PARALLELOGRAM, "trapezoid": MSO_SHAPE.TRAPEZOID,
            "pentagon": MSO_SHAPE.REGULAR_PENTAGON, "hexagon": MSO_SHAPE.HEXAGON, "octagon": MSO_SHAPE.OCTAGON,
            "star": MSO_SHAPE.STAR_5_POINT, "plus": MSO_SHAPE.CROSS, "chevron": MSO_SHAPE.CHEVRON,
            "callout": MSO_SHAPE.RECTANGULAR_CALLOUT, "arrow-right": MSO_SHAPE.RIGHT_ARROW,
        }
        elements = [
            {"type": "shape", "shapeType": kind, "x": 10 + i * 60, "y": 10, "width": "50px", "height": "50px", "styles": {"backgroundColor": "#2563eb"}}
            for i, kind in enumerate(expected)
        ]
        shapes = [s for s in self._slide(elements).shapes if s.shape_type == 1]
        self.assertEqual([s.auto_shape_type for s in shapes], list(expected.values()))

    def test_shape_text_goes_into_the_shapes_own_text_frame(self):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

        slide = self._slide([{
            "type": "shape", "shapeType": "rectangle", "x": 0, "y": 0, "width": "300px", "height": "150px",
            "styles": {"backgroundColor": "#1d4ed8"}, "shapeText": "Native state\n(folded)",
            "shapeTextStyle": {"color": "#ffffff", "fontSize": "24px", "fontWeight": "700", "textAlign": "center", "verticalAlign": "middle", "fontFamily": '"Manrope", sans-serif'},
        }])
        shape = next(s for s in slide.shapes if s.shape_type == 1)
        frame = shape.text_frame
        self.assertEqual([p.text for p in frame.paragraphs], ["Native state", "(folded)"])
        self.assertEqual(frame.vertical_anchor, MSO_ANCHOR.MIDDLE)
        run = frame.paragraphs[0].runs[0]
        self.assertEqual((frame.paragraphs[0].alignment, run.font.size.pt, run.font.bold, str(run.font.color.rgb), run.font.name), (PP_ALIGN.CENTER, 18.0, True, "FFFFFF", "Manrope"))

    def test_text_in_a_star_sits_in_the_editors_text_area_not_the_presets_small_one(self):
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        slide = self._slide([{
            "type": "shape", "shapeType": "star", "x": 96, "y": 96, "width": "192px", "height": "192px",
            "styles": {"backgroundColor": "#6d7fff"}, "shapeText": "Native",
            "shapeTextStyle": {"color": "#ffffff", "fontSize": "20px", "fontWeight": "600", "textAlign": "center", "verticalAlign": "middle"},
        }])
        group = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
        star = next(s for s in group.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
        box = next(s for s in group.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX)
        self.assertEqual(star.text_frame.text, "")
        self.assertEqual(box.text_frame.text, "Native")
        to_px = lambda emu: round(emu / 9525)
        # The editor's star text area: 26 % in from the sides, 30 % from the top, 22 % from the bottom.
        self.assertEqual((to_px(box.left), to_px(box.top), to_px(box.width), to_px(box.height)), (146, 154, 92, 92))
        self.assertTrue(box.text_frame.word_wrap)

    def test_table_fills_its_box_and_has_grid_lines(self):
        from pptx.util import Inches

        slide = self._slide([{
            "type": "table", "x": 0, "y": 0, "width": "480px", "height": "240px", "styles": {},
            "tableData": {"rows": 2, "cols": 2, "cells": [["A", "B"], ["1", "2"]], "colWidths": [100, 300], "rowHeights": [44, 44]},
        }])
        table = next(s for s in slide.shapes if s.has_table).table
        self.assertEqual([round(c.width / Inches(1), 2) for c in table.columns], [1.25, 3.75])
        self.assertEqual([round(r.height / Inches(1), 2) for r in table.rows], [1.25, 1.25])
        borders = table.cell(0, 0)._tc.tcPr.xpath("./a:lnL | ./a:lnR | ./a:lnT | ./a:lnB")
        self.assertEqual(len(borders), 4)

    def test_chart_without_a_card_on_a_dark_theme_has_light_labels(self):
        chart = {
            "type": "chart", "chartType": "bar", "x": 0, "y": 0, "width": "400px", "height": "300px", "styles": {},
            "chartData": {"labels": ["a", "b"], "datasets": [{"label": "s", "data": [1, 2]}]},
        }
        colours = {}
        for theme in ("horizon", "editorial"):
            shape = next(s for s in self._slide([chart], presentationTheme=theme).shapes if s.has_chart)
            rgb = shape.chart.font.color.rgb
            colours[theme] = sum(rgb) / 3
        self.assertGreater(colours["horizon"], 150)
        self.assertLess(colours["editorial"], 110)


class PPTXExportMediaTests(TestCase):
    def _export(self, elements):
        from pptx import Presentation
        from slideforge.bridge.pptx_exporter import PPTXExporter

        state = {"slides": [{"masterId": "none", "elements": elements}]}
        return Presentation(PPTXExporter(state).export()).slides[0]

    def _data_url(self, fmt, mime):
        import base64
        import io as _io
        from PIL import Image

        buffer = _io.BytesIO()
        Image.new("RGB", (40, 30), (220, 38, 38)).save(buffer, fmt)
        return f"data:{mime};base64," + base64.b64encode(buffer.getvalue()).decode()

    def test_bar_chart_exports_as_a_native_chart_with_its_data(self):
        slide = self._export([{
            "type": "chart", "chartType": "bar", "x": 100, "y": 100, "width": "500px", "height": "300px",
            "chartData": {"labels": ["Q1", "Q2", "Q3"], "datasets": [{"label": "Revenue", "data": [4.1, 5.3, "6"], "backgroundColor": "#2563eb"}]},
            "styles": {"backgroundColor": "#ffffff", "padding": "16px", "borderRadius": "12px"},
        }])
        charts = [shape.chart for shape in slide.shapes if shape.has_chart]
        self.assertEqual(len(charts), 1)
        plot = charts[0].plots[0]
        self.assertEqual(list(plot.categories), ["Q1", "Q2", "Q3"])
        self.assertEqual(list(plot.series[0].values), [4.1, 5.3, 6.0])
        self.assertEqual(plot.series[0].name, "Revenue")
        self.assertEqual(str(plot.series[0].format.fill.fore_color.rgb), "2563EB")
        self.assertFalse(any(shape.has_text_frame and "placeholder" in shape.text_frame.text for shape in slide.shapes))

    def test_one_series_chart_has_no_automatic_title_beside_its_legend(self):
        slide = self._export([{
            "type": "chart", "chartType": "bar", "x": 100, "y": 100, "width": "500px", "height": "300px",
            "chartData": {"labels": ["WT", "K12A"], "datasets": [{"label": "RMSD (Å)", "data": [2.1, 3.4]}]},
            "chartOptions": {"plugins": {"legend": {"display": True, "position": "top"}}},
        }])
        chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
        self.assertFalse(chart.has_title)
        self.assertTrue(chart.has_legend)
        titled = self._export([{
            "type": "chart", "chartType": "bar", "x": 100, "y": 100, "width": "500px", "height": "300px",
            "chartData": {"labels": ["WT"], "datasets": [{"label": "RMSD", "data": [2.1]}]},
            "chartOptions": {"plugins": {"title": {"display": True, "text": "Mutant stability"}}},
        }])
        chart = next(shape.chart for shape in titled.shapes if shape.has_chart)
        self.assertTrue(chart.has_title)
        self.assertEqual(chart.chart_title.text_frame.text, "Mutant stability")

    def test_chart_style_series_titles_and_legend_are_exported(self):
        from pptx.enum.chart import XL_LEGEND_POSITION

        slide = self._export([{
            "type": "chart", "chartType": "bar", "x": 0, "y": 0, "width": "500px", "height": "300px",
            "chartData": {"labels": ["WT", "K12A"], "datasets": [
                {"label": "Run 1", "data": [2.1, 3.4], "backgroundColor": "#4f7cff"},
                {"label": "Run 2", "data": [2.3, 3.1], "backgroundColor": "#dc2626"},
            ]},
            "chartOptions": {"plugins": {"legend": {"display": True, "position": "top"}}},
            "chartStyle": {"legend": "bottom", "title": "RMSD by mutant", "xTitle": "Mutant", "yTitle": "RMSD (Å)", "fontSize": 16, "grid": False},
            "styles": {"backgroundColor": "#ffffff", "padding": "16px"},
        }])
        chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
        series = chart.plots[0].series
        self.assertEqual([s.name for s in series], ["Run 1", "Run 2"])
        self.assertEqual(str(series[1].format.fill.fore_color.rgb), "DC2626")
        self.assertEqual(chart.chart_title.text_frame.text, "RMSD by mutant")
        self.assertEqual(chart.legend.position, XL_LEGEND_POSITION.BOTTOM)
        self.assertEqual(chart.category_axis.axis_title.text_frame.text, "Mutant")
        self.assertEqual(chart.value_axis.axis_title.text_frame.text, "RMSD (Å)")
        self.assertFalse(chart.value_axis.has_major_gridlines)
        self.assertEqual(chart.font.size.pt, 12.0)
        hidden = self._export([{
            "type": "chart", "chartType": "pie", "x": 0, "y": 0, "width": "300px", "height": "300px",
            "chartData": {"labels": ["a", "b"], "datasets": [{"label": "s", "data": [1, 2]}]},
            "chartStyle": {"legend": "none"},
        }])
        self.assertFalse(next(shape.chart for shape in hidden.shapes if shape.has_chart).has_legend)

    def test_connector_follows_its_points_and_keeps_its_arrowhead(self):
        from pptx.util import Emu

        slide = self._export([{
            "type": "connector", "connectorType": "line", "connectorEnd": "arrow", "connectorStart": "none",
            "x": 100, "y": 100, "width": "300px", "height": "200px",
            "points": [{"x": 28, "y": 172}, {"x": 272, "y": 28}],  # rising left to right, inside the padded box
            "styles": {"color": "#0e7490", "strokeWidth": 4},
        }, {
            "type": "connector", "connectorType": "poly", "x": 0, "y": 400, "width": "300px", "height": "200px",
            "points": [{"x": 28, "y": 28}, {"x": 150, "y": 28}, {"x": 150, "y": 172}],
            "styles": {"color": "#0e7490", "strokeWidth": 4},
        }])
        line = next(s for s in slide.shapes if s.shape_type == 6 or getattr(s, "begin_x", None) is not None)
        to_px = lambda emu: round(Emu(emu).inches * 96)
        self.assertEqual((to_px(line.begin_x), to_px(line.begin_y)), (128, 272))
        self.assertEqual((to_px(line.end_x), to_px(line.end_y)), (372, 128))
        ends = {child.tag.split("}")[1]: child.get("type") for child in line.line._get_or_add_ln()}
        self.assertEqual(ends.get("tailEnd"), "triangle")
        self.assertNotIn("headEnd", ends)
        elbow = [s for s in slide.shapes if s.shape_type == 5]  # a freeform with the bend
        self.assertEqual(len(elbow), 1)

    def test_line_chart_exports_as_a_line_chart(self):
        from pptx.enum.chart import XL_CHART_TYPE

        slide = self._export([{
            "type": "chart", "chartType": "line", "x": 0, "y": 0, "width": "400px", "height": "300px",
            "chartData": {"labels": ["a", "b"], "datasets": [{"label": "s", "data": [1, 2], "borderColor": "#16a34a"}]},
            "styles": {},
        }])
        chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
        self.assertEqual(chart.chart_type, XL_CHART_TYPE.LINE_MARKERS)

    def test_webp_photo_is_embedded_as_png(self):
        slide = self._export([{"type": "image", "content": self._data_url("WEBP", "image/webp"), "x": 0, "y": 0, "width": "400px", "height": "300px", "styles": {}}])
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        self.assertEqual(len(pictures), 1)
        self.assertEqual(pictures[0].image.content_type, "image/png")

    def test_svg_picture_is_embedded(self):
        import base64

        svg = '<svg xmlns="http://www.w3.org/2000/svg" width="40" height="30"><rect width="40" height="30" fill="#0ea5e9"/></svg>'
        content = "data:image/svg+xml;base64," + base64.b64encode(svg.encode()).decode()
        slide = self._export([{"type": "image", "content": content, "x": 0, "y": 0, "width": "400px", "height": "300px", "styles": {}}])
        self.assertEqual(len([shape for shape in slide.shapes if shape.shape_type == 13]), 1)
