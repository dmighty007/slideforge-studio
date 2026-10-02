import io
import json

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image
from pptx import Presentation as PptxPresentation
from django.test import Client, TestCase, override_settings

from slideforge.desktop.middleware import get_local_user, launch_cookie_name
from slideforge.desktop.paths import FRONTEND_DIR


class LocalUserTests(TestCase):
    def test_session_reports_local_user_without_sign_in(self):
        response = Client().get("/api/auth/session/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.json(),
            {"authenticated": True, "user": {"id": get_local_user().id, "username": "local"}, "accountsEnabled": False},
        )

    def test_local_user_cannot_sign_in_with_a_password(self):
        self.assertFalse(get_local_user().has_usable_password())

    def test_sign_in_endpoints_are_gone(self):
        client = Client()
        for path in ("/api/auth/login/", "/api/auth/register/", "/api/auth/logout/", "/admin/"):
            self.assertEqual(client.post(path).status_code, 404, path)

    def test_presentations_are_created_for_local_user(self):
        response = Client().post(
            "/api/presentations/",
            data='{"title": "Deck", "state": {"slides": []}}',
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()["title"], "Deck")


@override_settings(SLIDEFORGE_LAUNCH_TOKEN="launch-secret")
class LaunchTokenTests(TestCase):
    def test_request_without_token_is_refused(self):
        self.assertEqual(Client().get("/").status_code, 403)
        self.assertEqual(Client().get("/api/auth/session/").status_code, 403)

    def test_bundled_frontend_code_loads_without_the_cookie(self):
        # The sandboxed molecule viewer fetches vendor/ngl/ngl.js without cookies.
        for path in ("/vendor/dompurify/purify.min.js", "/js/core/state.js"):
            response = Client().get(path)
            self.assertEqual(response.status_code, 200, path)
            response.close()
        for path in ("/", "/api/auth/session/", "/media/assets/x.png"):
            self.assertEqual(Client().get(path).status_code, 403, path)

    def test_wrong_token_is_refused(self):
        self.assertEqual(Client().get("/?token=nope").status_code, 403)
        self.assertEqual(Client().get("/?token=%C3%A9").status_code, 403)

    def test_token_is_traded_for_cookie_and_removed_from_url(self):
        client = Client()
        response = client.get("/?token=launch-secret&slide=2")

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], "/?slide=2")
        cookie = response.cookies[launch_cookie_name(80)]
        self.assertEqual(cookie.value, "launch-secret")
        self.assertTrue(cookie["httponly"])
        self.assertEqual(client.get("/").status_code, 200)
        self.assertEqual(client.get("/api/auth/session/").status_code, 200)


class DesktopServingTests(TestCase):
    def test_frontend_assets_are_served_with_debug_off(self):
        self.assertFalse(settings.DEBUG)
        client = Client()
        for path in ("/", "/js/core/state.js", "/css/base.css", "/vendor/dompurify/purify.min.js"):
            response = client.get(path)
            self.assertEqual(response.status_code, 200, path)
            if hasattr(response, "close"):
                response.close()

    def test_csp_only_on_html_so_the_pdf_viewer_is_not_blocked(self):
        (settings.MEDIA_ROOT / "csp-test.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
        client = Client()
        page = client.get("/")
        pdf = client.get("/media/csp-test.pdf")

        self.assertIn("Content-Security-Policy", page)
        self.assertEqual(pdf.status_code, 200)
        self.assertNotIn("Content-Security-Policy", pdf)
        for response in (page, pdf):
            response.close()

    def test_data_lives_in_the_data_directory(self):
        # The database path is checked end to end by the installed-app test (tests run on a test database).
        self.assertTrue(str(settings.MEDIA_ROOT).startswith(str(settings.DATA_DIR)))
        self.assertTrue((settings.DATA_DIR / "secret_key").exists())


class NoSignInUiTests(TestCase):
    def test_page_is_served_with_accounts_disabled(self):
        response = Client().get("/")
        self.assertContains(response, '<body class="accounts-disabled ')

    def test_css_hides_all_sign_in_ui_when_accounts_are_disabled(self):
        css = (FRONTEND_DIR / "css" / "base.css").read_text(encoding="utf-8")
        for selector in ("body.accounts-disabled #entry-hero", "body.accounts-disabled #auth-modal", "body.accounts-disabled .toolbar-auth"):
            self.assertIn(selector, css)

    def test_session_reports_accounts_disabled(self):
        self.assertFalse(Client().get("/api/auth/session/").json()["accountsEnabled"])


class UploadedMediaExportTests(TestCase):
    def test_uploaded_image_is_included_in_pptx_export(self):
        # Uploads live in the data directory, not the package; the exporter must still find them.
        png = io.BytesIO()
        Image.new("RGB", (64, 48), (200, 30, 30)).save(png, format="PNG")
        client = Client()
        upload = client.post("/api/assets/upload/", {"file": SimpleUploadedFile("red.png", png.getvalue(), "image/png")})
        self.assertEqual(upload.status_code, 201, upload.content)
        image = {"id": "el_1", "type": "image", "content": upload.json()["url"], "x": 10, "y": 10,
                 "width": "200px", "height": "150px", "styles": {}}
        state = {"slides": [{"id": "slide_1", "elements": [image]}]}

        response = client.post(
            "/api/presentations/export/pptx/",
            data=json.dumps({"state": state, "title": "Deck"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        deck = PptxPresentation(io.BytesIO(b"".join(response.streaming_content)))
        pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13]
        self.assertEqual(len(pictures), 1)


class FrontendCachingTests(TestCase):
    """After an update the window must not mix cached old scripts with new ones (ES modules import each other by
    plain relative paths, so a stale module breaks the page)."""

    def test_scripts_styles_and_the_page_are_always_revalidated(self):
        client = Client()
        for path in ("/js/mermaid/mermaid-graph.js", "/css/base.css", "/"):
            response = client.get(path)
            self.assertEqual(response.status_code, 200, path)
            self.assertEqual(response["Cache-Control"], "no-cache", path)


class ShareTunnelTests(TestCase):
    """The public tunnel reaches shared presentations (/s/...) and nothing else."""

    def call(self, path, method="GET"):
        from slideforge.desktop.share_tunnel import share_only_app

        captured = {}

        def start_response(status, headers, exc_info=None):
            captured["status"] = int(status.split()[0])

        environ = {
            "REQUEST_METHOD": method, "PATH_INFO": path, "QUERY_STRING": "", "SERVER_NAME": "x.trycloudflare.com",
            "SERVER_PORT": "443", "HTTP_HOST": "x.trycloudflare.com", "wsgi.url_scheme": "https",
            "wsgi.input": io.BytesIO(), "wsgi.errors": io.StringIO(), "SERVER_PROTOCOL": "HTTP/1.1",
        }
        body = share_only_app(environ, start_response)
        b"".join(body)
        if hasattr(body, "close"):
            body.close()
        return captured["status"]

    def test_editor_api_and_media_are_not_reachable_through_the_tunnel(self):
        for path in ("/", "/index.html", "/api/presentations/", "/api/share/tunnel/", "/media/x.png", "/js/core/state.js"):
            self.assertEqual(self.call(path), 404, path)
        self.assertEqual(self.call("/s/AAAAAAAAAAAAAAAAAAAAAA/", method="POST"), 404)

    def test_shared_presentation_is_served_through_the_tunnel(self):
        from slideforge.studio import sharing
        from slideforge.studio.models import Presentation, ShareLink

        deck = Presentation.objects.create(owner=get_local_user(), title="Deck")
        share = ShareLink.objects.create(presentation=deck, token=sharing.new_token())
        self.addCleanup(sharing.delete_bundle, share.token)
        buffer = io.BytesIO()
        import zipfile

        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("index.html", "<!doctype html><title>Deck</title>")
        sharing.publish(share, buffer.getvalue())

        self.assertEqual(self.call(f"/s/{share.token}/"), 200)
        self.assertEqual(self.call("/s/AAAAAAAAAAAAAAAAAAAAAA/"), 404)

    @override_settings(SLIDEFORGE_LAUNCH_TOKEN="launch-secret")
    def test_tunnel_controls_need_the_launch_token(self):
        self.assertEqual(Client().get("/api/share/tunnel/").status_code, 403)
        self.assertEqual(Client().post("/api/share/tunnel/", data='{"action": "start"}', content_type="application/json").status_code, 403)

    def test_tunnel_address_is_read_from_cloudflared(self):
        import os
        import sys
        import tempfile
        from pathlib import Path
        from unittest import mock

        from slideforge.desktop.share_tunnel import ShareTunnel

        with tempfile.TemporaryDirectory() as data_dir:
            fake = Path(data_dir) / "bin" / "cloudflared"
            fake.parent.mkdir()
            # Prints what cloudflared prints, records its arguments, then stays up like a running tunnel.
            fake.write_text(
                f"#!{sys.executable}\nimport sys, time, pathlib\n"
                f"pathlib.Path({str(Path(data_dir) / 'args')!r}).write_text(' '.join(sys.argv[1:]))\n"
                "print('INF |  https://quiet-river-demo.trycloudflare.com                                  |', flush=True)\n"
                "time.sleep(60)\n"
            )
            fake.chmod(0o755)
            tunnel = ShareTunnel(Path(data_dir))
            with mock.patch("shutil.which", return_value=None):
                self.assertEqual(tunnel.binary(), str(fake))
                tunnel.start()
                try:
                    self.assertEqual(tunnel.wait_for_url(15), "https://quiet-river-demo.trycloudflare.com")
                    status = tunnel.status()
                    self.assertTrue(status["running"])
                    self.assertEqual(status["url"], "https://quiet-river-demo.trycloudflare.com")
                    args = (Path(data_dir) / "args").read_text()
                    port = tunnel._server.effective_port
                    self.assertIn(f"--url http://127.0.0.1:{port}", args)
                    self.assertIn("--no-autoupdate", args)
                finally:
                    tunnel.shutdown()
                self.assertFalse(tunnel.status()["running"])
                self.assertIsNone(tunnel.status()["url"])


class QtPermissionTests(TestCase):
    """pywebview's own permission handler passes an int where PyQt6 wants an enum; the TypeError aborted the app
    the first time a page asked for the clipboard (Share > Copy)."""

    def test_permission_requests_are_answered_with_enums_and_never_raise(self):
        import unittest

        from slideforge.desktop import qt_support

        if not qt_support.available():
            raise unittest.SkipTest("Qt WebEngine is not installed")
        from qtpy.QtCore import QUrl
        from qtpy.QtWebEngineCore import QWebEnginePage
        from webview.platforms import qt as pywebview_qt

        qt_support._app_port = 8123
        qt_support._patch_permissions()
        handler = pywebview_qt.BrowserView.WebPage.onFeaturePermissionRequested
        answers = []

        class Page:
            def setFeaturePermission(self, url, feature, policy):
                assert isinstance(policy, QWebEnginePage.PermissionPolicy), type(policy)
                answers.append((url.port(), feature, policy))

        feature = QWebEnginePage.Feature
        policy = QWebEnginePage.PermissionPolicy
        handler(Page(), QUrl("http://127.0.0.1:8123"), feature.ClipboardReadWrite)
        handler(Page(), QUrl("https://example.com"), feature.ClipboardReadWrite)
        handler(Page(), QUrl("http://127.0.0.1:8123"), feature.Geolocation)
        self.assertEqual(
            [answer[2] for answer in answers],
            [policy.PermissionGrantedByUser, policy.PermissionDeniedByUser, policy.PermissionDeniedByUser],
        )

        class BrokenPage:
            def setFeaturePermission(self, *args):
                raise TypeError("boom")

        handler(BrokenPage(), QUrl("http://127.0.0.1:8123"), feature.ClipboardReadWrite)  # logged, not raised


class QtPopupNavigationTests(TestCase):
    """Opening a share link (an outside address) from the app closed its popup inside Qt's navigation callback;
    Qt WebEngine then aborted the whole app with SIGTRAP."""

    def test_outside_link_goes_to_the_browser_and_the_popup_closes_only_afterwards(self):
        from unittest import mock

        from slideforge.desktop import qt_support

        class Url:
            def __init__(self, scheme, host, port, text):
                self._parts = (scheme, host, port, text)

            def scheme(self):
                return self._parts[0]

            def host(self):
                return self._parts[1]

            def port(self):
                return self._parts[2]

            def toString(self):
                return self._parts[3]

        class View:
            closed = 0

            def close(self):
                self.closed += 1

        qt_support._app_port = 8123
        outside = Url("https", "demo.trycloudflare.com", -1, "https://demo.trycloudflare.com/s/abc/")
        inside = Url("http", "127.0.0.1", 8123, "http://127.0.0.1:8123/presenter")
        view, scheduled = View(), []
        with mock.patch.object(qt_support.webbrowser, "open") as browser:
            self.assertFalse(qt_support._route_popup_navigation(view, outside, True, scheduled.append))
            browser.assert_called_once_with("https://demo.trycloudflare.com/s/abc/", 2, True)
            self.assertEqual(view.closed, 0, "the popup was closed inside the navigation callback")
            self.assertEqual(len(scheduled), 1)
            scheduled[0]()
            self.assertEqual(view.closed, 1)

            # The app's own pages (presenter view) and sub-frames load in the popup.
            self.assertTrue(qt_support._route_popup_navigation(view, inside, True, scheduled.append))
            self.assertTrue(qt_support._route_popup_navigation(view, outside, False, scheduled.append))
            self.assertEqual((view.closed, len(scheduled)), (1, 1))

            # A browser that cannot be started must not raise out of the Qt callback.
            browser.side_effect = OSError("no browser")
            self.assertFalse(qt_support._route_popup_navigation(view, outside, True, scheduled.append))


class QtTerminateTests(TestCase):
    """`kill` on the desktop window raised SystemExit inside a Qt callback: the process aborted and nothing was
    cleaned up (instance.json left behind, a share tunnel left running)."""

    def test_sigterm_asks_qt_to_quit_and_never_raises(self):
        import signal
        import unittest
        from unittest import mock

        from slideforge.desktop import qt_support

        if not qt_support.available():
            raise unittest.SkipTest("Qt WebEngine is not installed")
        previous = signal.getsignal(signal.SIGTERM)
        try:
            qt_support._quit_on_terminate()
            handler = signal.getsignal(signal.SIGTERM)
            self.assertIs(handler, qt_support._request_quit)
            with mock.patch("qtpy.QtWidgets.QApplication") as application:
                handler(signal.SIGTERM, None)  # must not raise SystemExit
                application.instance.return_value.quit.assert_called_once_with()
                application.instance.side_effect = RuntimeError("no app")
                handler(signal.SIGTERM, None)  # logged, not raised
        finally:
            signal.signal(signal.SIGTERM, previous)


class QtFullscreenTests(TestCase):
    """Leaving a presentation puts the window back as it was: a maximized window came back small."""

    def test_window_returns_to_its_state_before_fullscreen(self):
        from slideforge.desktop import qt_support

        class Signal:
            def connect(self, handler):
                self.handler = handler

        class Page:
            fullScreenRequested = Signal()

        class Window:
            def __init__(self, maximized):
                self.maximized, self.fullscreen, self.calls = maximized, False, []

            def isFullScreen(self):
                return self.fullscreen

            def isMaximized(self):
                return self.maximized

            def showFullScreen(self):
                self.fullscreen = True
                self.calls.append("fullscreen")

            def showMaximized(self):
                self.fullscreen = False
                self.calls.append("maximized")

            def showNormal(self):
                self.fullscreen = False
                self.calls.append("normal")

        class Request:
            def __init__(self, on):
                self.on = on

            def accept(self):
                pass

            def toggleOn(self):
                return self.on

        for maximized, expected in ((True, "maximized"), (False, "normal")):
            page, window = Page(), Window(maximized)
            qt_support._follow_fullscreen(page, window)
            for _ in range(2):  # twice: the remembered state must survive a second presentation
                page.fullScreenRequested.handler(Request(True))
                page.fullScreenRequested.handler(Request(False))
            self.assertEqual(window.calls, ["fullscreen", expected, "fullscreen", expected])
